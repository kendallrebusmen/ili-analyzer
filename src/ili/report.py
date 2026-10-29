"""Report (spec 6.10): the self-contained HTML report, and the same report as a PDF deliverable (``report_pdf``).

``render_report`` draws every section from the report inputs (``report_data``), the category changes
(``category_change``) and the charts (``report_charts``) as they are; nothing is recomputed. The page is
fully local: plotly.js is inlined once and every chart is a div without its own copy, and nothing is
fetched. jinja2 escapes every value; only the plotly.js bundle and the chart divs are marked safe, so a
vendor id, a file path or the line name can never become markup. None / NaN render as an empty cell.

The report is always DRAFT: every section opens with the label, a banner carries it on screen and a page
margin box repeats it on each printed page,
and no wording marks it final or approved.

DRAFT – for review by qualified personnel.
"""

from __future__ import annotations

import argparse
import json
import os
import subprocess
import sys
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path

import jinja2
import numpy as np
import pandas as pd
import plotly.offline
import yaml
from markupsafe import Markup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ili import __version__
from ili.align_welds import DRAFT_LABEL
from ili.analysis_table import (ANALYSIS_TABLE_FILE, ANALYSIS_XLSX_FILE, SHEET_NAME, build_analysis_table)
from ili.brand import product_label
from ili.category_change import CHANGE_COLUMNS, CategoryChangeResult, category_changes
from ili.classify import CATEGORY_PRIORITY, IMMEDIATE_ACTION_NOTE, REMOVED_CATEGORY, criteria_path
from ili.config import DEFAULT_SETTINGS, PROJECT_ROOT, Settings, load_settings
from ili.match import CONFIDENCES
from ili.report_charts import (CATEGORY_COLOURS, depth_vs_distance, elevation_chart, feature_map, figure_div,
                               growth_histograms, growth_vs_distance, orientation_charts, rl_distribution,
                               unity_chart)
from ili.report_data import (ReportInputs, _sha256, backtest, config_rows, criteria_rows, dig_list, input_hashes,
                             largest_growth, load_report_inputs, match_stats, pending_reviews, removed_features)
from ili.report_pdf import build_pdf
from ili.report_tables import (DEPTH_COLUMNS, LOCATIONS, anomaly_distribution, anomaly_distribution_table,
                               anomalies_of_interest, imp_class_rows, imp_class_summary, joint_ml_counts,
                               largest_joints, length_check, pipe_summary)
from ili.review import FLAG_COLUMNS, REVIEW_COLUMNS

TEMPLATES = Path(__file__).parent / "templates"
TEMPLATE_NAME = "report.html.j2"
SECTION_IDS = ("header", "summary", "data-quality", "growth", "dig-list", "largest-changes", "charts", "methods")
NOT_STATED = "not stated"
DECIMALS = 3  # presentation only: the CSV outputs keep full precision
COORD_DECIMALS = 6
COORD_COLUMNS = ("lat", "lon")
CSS_UNSAFE = "\"\\'<>&\n"  # changed by CSS string syntax or by HTML autoescaping inside <style>

IMMEDIATE_COLUMNS = ["match_id", "prev_feature_id", "curr_feature_id", "feature_type", "curr_log_dist_ft",
                     "governing_rule", "triggered_rules", "action_note"]
LARGEST_GROWTH_COLUMNS = ["match_id", "match_type", "prev_feature_id", "curr_feature_id", "curr_log_dist_ft",
                          "segment_id", "history_class", "d_prev_pct", "d_curr_pct", "raw_mpy", "governing_mpy",
                          "governing_method", "needs_review"]
MOVED_COLUMNS = [c for c in CHANGE_COLUMNS if c != "moved_more_severe"]
# Printed (landscape) the wide tables keep only these columns; dig_list.xlsx and category_changes.csv hold all.
DIG_LIST_PRINT_COLUMNS = ("governing_category", "governing_rule", "match_id", "curr_feature_id", "curr_log_dist_ft",
                          "curr_us_gw_id", "curr_dist_from_us_gw_ft", "clock_min", "depth_pct_assessed", "rl_yr",
                          "due_date", "requires_review")
MOVED_PRINT_COLUMNS = ("match_id", "curr_feature_id", "feature_type", "prev_depth_pct_assessed",
                       "curr_depth_pct_assessed", "prev_rule_category", "prev_rules_not_evaluable",
                       "curr_rule_category", "curr_governing_rule", "curr_rules_not_evaluable")
UNRECORDED = "saved after the last run — rerun to apply"
REVIEW_LOG_COLUMNS = [*REVIEW_COLUMNS, "applied_status", "reason"]  # reviews.csv joined to review_status.csv
FLAG_LOG_COLUMNS = [*FLAG_COLUMNS, "applied_status", "reason"]  # feature_flags.csv with what became of each row


@dataclass(frozen=True)
class ReportHeader:
    """Header fields the tallies do not carry (CLI options); None renders as "not stated"."""
    line_name: str | None
    prev_tool: str | None
    curr_tool: str | None
    prepared_for: str | None = None  # the PDF cover's client (operator) name
    prepared_by: str | None = None  # the PDF cover's author (company or engineer)


@dataclass(frozen=True)
class Table:
    """A table ready for the template: every cell already a string ("" for a blank). With ``print_columns``
    the other columns are hidden in print."""
    columns: list[str]
    rows: list[list[str]]
    row_classes: list[str] | None = None
    print_columns: tuple[str, ...] | None = None

    @property
    def hidden_in_print(self) -> set[int]:
        """Indices of the columns not in ``print_columns`` (none without it)."""
        if self.print_columns is None:
            return set()
        return {i for i, c in enumerate(self.columns) if c not in self.print_columns}


def _number(value: float, decimals: int) -> str:
    text = f"{value:.{decimals}f}"
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return "0" if text == "-0" else text


def cell(value, column: str | None = None) -> str:
    """One value as display text: None, NaN, NaT and NA are ""; floats to ``DECIMALS`` places
    (``COORD_DECIMALS`` for lat/lon) without trailing zeros; dates as ISO; lists comma-joined."""
    if isinstance(value, (list, tuple)):
        return ", ".join(cell(v) for v in value)
    if value is None or (not isinstance(value, (str, dict)) and pd.isna(value)):
        return ""
    if isinstance(value, (bool, np.bool_)):
        return str(bool(value))
    if isinstance(value, (int, np.integer)):
        return str(int(value))
    if isinstance(value, (float, np.floating)):
        return _number(float(value), COORD_DECIMALS if column in COORD_COLUMNS else DECIMALS)
    if isinstance(value, pd.Timestamp):
        return value.date().isoformat() if value == value.normalize() else value.isoformat()
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def table(df: pd.DataFrame, columns: list[str] | None = None, row_classes: list[str] | None = None,
          print_columns: tuple[str, ...] | None = None) -> Table:
    """``df[columns]`` as a ``Table``; ``row_classes`` gives each row's class ("" for none)."""
    columns = list(df.columns) if columns is None else list(columns)
    rows = [[cell(v, c) for c, v in zip(columns, r, strict=True)] for r in df[columns].itertuples(index=False)]
    return Table(columns, rows, row_classes, print_columns)


def _css_string(text: str) -> str:
    """``text`` for a double-quoted CSS string in an autoescaped ``<style>``: refused if it holds a character
    that CSS or the HTML escaping would change."""
    unsafe = sorted(set(text) & set(CSS_UNSAFE))
    if unsafe:
        raise ValueError(f"label {text!r} contains {unsafe}, which cannot go in the print margin box")
    return text


def _counts(mapping: dict, key: str) -> Table:
    return Table([key, "count"], [[cell(k), cell(v)] for k, v in mapping.items()])


def _vendor(manifest: dict, run_id: str) -> str:
    """The ``vendor`` key of the run's vendor map named in the ingest manifest; "not stated" if unreadable. A
    relative path not found from the working directory is read from the project root (ingest records it as
    given, usually relative to the project)."""
    try:
        run = next(r for r in manifest["runs"] if r["run_id"] == run_id)
        path = Path(run["vendor_map"])
        if not path.is_absolute() and not path.is_file():
            path = PROJECT_ROOT / path
        vmap = yaml.safe_load(path.read_text(encoding="utf-8"))
        vendor = vmap["vendor"]
    except (KeyError, TypeError, StopIteration, OSError, yaml.YAMLError):
        return NOT_STATED
    return cell(vendor) or NOT_STATED


def _run_dates(std: pd.DataFrame) -> str:
    return ", ".join(sorted({cell(d) for d in std.run_date.dropna()})) or NOT_STATED


def _rate(value) -> str:
    return "not available" if value is None else f"{100.0 * float(value):.2f} %"


def _header(inputs: ReportInputs, header: ReportHeader, generated_utc: datetime) -> dict:
    match = inputs.summaries["match"]
    prev_id, curr_id = match["prev_run_id"], match["curr_run_id"]
    runs = Table(["run", "run id", "run date", "vendor", "tool type"], [
        [role, run_id, _run_dates(std), _vendor(inputs.ingest_manifest, run_id), tool or NOT_STATED]
        for role, run_id, std, tool in (("previous", prev_id, inputs.prev_std, header.prev_tool),
                                        ("current", curr_id, inputs.curr_std, header.curr_tool))])
    ps = pipe_summary(inputs.pipe, inputs.mop, inputs.segments, inputs.settings)

    def span(pair) -> str:
        return "not available" if pair is None else f"{cell(pair[0])} to {cell(pair[1])}"

    pipe = {"by_wall_grade": table(ps.by_wall_grade),
            "design_pressure": span(ps.design_pressure_psig), "stress": span(ps.stress_at_mop_pct_smys),
            "install_years": cell(ps.install_years) or NOT_STATED, "coatings": cell(ps.coatings) or NOT_STATED,
            "hca_length_ft": cell(ps.hca_length_ft),
            "hca_pct": "not available" if ps.hca_pct is None else cell(ps.hca_pct),
            "design_factor": cell(inputs.settings.strength.design_factor),
            "length_check": table(length_check(inputs.prev_std, inputs.curr_std, inputs.pipe))}
    return {"line_name": header.line_name or NOT_STATED, "prev_run_id": prev_id, "curr_run_id": curr_id,
            "runs": runs, "pipe": pipe, "generated_utc": generated_utc.isoformat(timespec="seconds"),
            "generator": f"ili.report {__version__}", "product": product_label(),
            "prepared_for": header.prepared_for,
            "prepared_by": header.prepared_by}


def _imp_class_rows(inputs: ReportInputs) -> pd.DataFrame:
    """Each classified feature's IMP class (``imp_class.csv``; the summary table is aggregated from it)."""
    return imp_class_rows(inputs.classification, inputs.match_table, inputs.remaining_life, inputs.repairs,
                          inputs.settings)


def _summary(inputs: ReportInputs, dig: pd.DataFrame) -> dict:
    cls = inputs.classification
    removed = cls.governing_category.eq(REMOVED_CATEGORY)
    shown = [*CATEGORY_PRIORITY, *([REMOVED_CATEGORY] if removed.any() else [])]
    counts = cls.governing_category.value_counts().reindex(shown, fill_value=0)
    rl = inputs.remaining_life
    removed_ids = set(cls.match_id[removed])
    return {
        "immediate": table(dig[dig.governing_category == "immediate"], IMMEDIATE_COLUMNS),
        "immediate_note": IMMEDIATE_ACTION_NOTE,
        "category_counts": Table(["category", "count"], [[c, cell(n)] for c, n in counts.items()],
                                 [f"cat-{c}" for c in counts.index]),
        "imp_class": table(imp_class_summary(_imp_class_rows(inputs), inputs.settings, inputs.repairs is not None)),
        "ytc_flag_years": cell(inputs.settings.report.ytc_flag_years),
        "repairs_given": inputs.repairs is not None,
        "n_features": len(cls),
        "n_rl_dig": int((rl.dig_required_by_remaining_life.eq(True) & ~rl.match_id.isin(removed_ids)).sum()),
        "n_removed": int(removed.sum()),
        "removed": table(removed_features(inputs)),
        "n_rl_feature": len(inputs.remaining_life),
        "n_imp_scheduled": int(counts["imp_scheduled"]),
        "intervals": table(inputs.intervals),
        "aoi": table(anomalies_of_interest(inputs)),
        "aoi_count": cell(inputs.settings.report.aoi_count),
        "discovery_warning": inputs.summaries["classification"].get("discovery_warning"),
    }


def _match_stats(match_table: pd.DataFrame) -> pd.DataFrame:
    """``match_stats`` with the confidence columns in band order (high, medium, low), then ``all``."""
    stats = match_stats(match_table)
    confidences = [c for c in CONFIDENCES if c in stats.columns]
    others = [c for c in stats.columns if c not in confidences and c != "all"]
    return stats[[*confidences, *others, "all"]].reset_index()


def _data_quality(inputs: ReportInputs) -> dict:
    issues = inputs.validation_issues
    by_reason = issues.groupby(["run_id", "reason"], dropna=False).size().rename("count").reset_index()
    align = inputs.summaries["alignment"]
    cls_summary = inputs.summaries["classification"]
    n_pending, n_unrecorded = len(pending_reviews(inputs.match_table)), _n_unrecorded(inputs)
    return {
        "n_issues": len(issues),
        "issue_counts": table(by_reason),
        "issues": table(issues),
        "alignment_rate": _rate(align.get("alignment_rate")),
        "alignment_rate_outside_repairs": _rate(align.get("alignment_rate_outside_repairs")),
        "match_stats": table(_match_stats(inputs.match_table)),
        "n_pending": n_pending,
        "n_unrecorded": n_unrecorded,
        "pending_text": pending_text(n_pending, n_unrecorded),
        "not_final": n_pending + n_unrecorded > 0,
        "n_requires_review": cls_summary.get("n_requires_review"),
        "review_reasons": _counts(cls_summary.get("counts_by_review_reason") or {}, "reason"),
        "crossings_given": inputs.crossings is not None,
        "classified_with_crossings": cls_summary.get("crossings_file_given"),
    }


STATS_SCOPE = "the segment and joint percentiles and these statistics"


def recoat_note(summary: dict, settings: Settings) -> str:
    """What the recoat exclusion left out of the percentiles and the per-segment statistics, from
    ``growth_summary.json`` (``repairs_file_given``, ``n_excluded_recoated``) and
    ``growth.exclude_recoated_from_stats`` (checked against the summary when it records it)."""
    if summary.get("repairs_file_given") is False:
        return f"No repairs file was given to the growth step, so no recoated feature was left out of {STATS_SCOPE}."
    if not settings.growth.exclude_recoated_from_stats:
        return (f"No recoated feature was left out of {STATS_SCOPE}, because exclude_recoated_from_stats is "
                "false.")
    n = summary.get("n_excluded_recoated")
    if n is None:
        return f"growth_summary.json does not record how many recoated features were left out of {STATS_SCOPE}."
    return f"{n} recoated feature(s) left out of {STATS_SCOPE} (exclude_recoated_from_stats: true)."


def _growth(inputs: ReportInputs) -> dict:
    summary = inputs.summaries["growth"]
    pct_negative = cell(summary.get("pct_negative_raw"))
    return {
        "segments": table(inputs.growth_segments),
        "profile": summary.get("governing_profile"),
        "recoat_note": recoat_note(summary, inputs.settings),
        "by_surface": table(inputs.growth_by_surface),
        "pct_negative_raw": (f"{pct_negative} % of matched features in the statistics (recoated left out when the "
                             "setting is on)") if pct_negative else "not available",
        "methods": _counts(summary.get("counts_by_governing_method") or {}, "governing_method"),
        "histograms": Markup(figure_div(growth_histograms(inputs))),
    }


LOCATION_TITLES = {"all": "all metal loss", "girth_weld": "metal loss at a girth weld",
                   "long_seam": "metal loss on the long seam"}


def _distribution(inputs: ReportInputs) -> dict:
    """The depth-bin tables (one per location, both runs) and the dent-bin table, with the notes."""
    result = anomaly_distribution(inputs)
    columns = [c for c in DEPTH_COLUMNS if c != "location"]
    return {
        "depth": [(LOCATION_TITLES[loc], table(result.depth[result.depth.location == loc], columns),
                   [n.text for n in result.notes if n.table == "depth" and n.location == loc])
                  for loc in LOCATIONS],
        "dent": table(result.dent),
        "dent_notes": [n.text for n in result.notes if n.table == "dent"],
    }


def _joints(inputs: ReportInputs) -> dict:
    """The per-joint metal-loss tables: the joints with the largest current count and the largest change."""
    n = inputs.settings.report.largest_changes_n
    by_count, by_change = largest_joints(joint_ml_counts(inputs.weld_map, inputs.prev_std, inputs.curr_std), n)
    return {"n": n, "by_count": table(by_count), "by_change": table(by_change)}


def _charts(inputs: ReportInputs, dig: pd.DataFrame) -> dict:
    tolerance = inputs.settings.tool.depth_tolerance_pct_wt
    bt = backtest(inputs.digs, inputs.curr_std, inputs.settings)
    fmap, unity, elevation = feature_map(inputs, dig), unity_chart(bt, tolerance), elevation_chart(inputs)
    return {
        "depth": Markup(figure_div(depth_vs_distance(inputs))),
        "growth": Markup(figure_div(growth_vs_distance(inputs))),
        "rl": Markup(figure_div(rl_distribution(inputs))),
        "orientation": [Markup(figure_div(fig)) for fig in orientation_charts(inputs).values()],
        "elevation": None if elevation is None else Markup(figure_div(elevation)),
        "map": None if fmap is None else Markup(figure_div(fmap)),
        "unity": None if unity is None else Markup(figure_div(unity)),
        "backtest": None if bt is None else {
            "n_within": bt.n_within, "n": len(bt.table), "pct_within": cell(bt.pct_within),
            "unmatched": ", ".join(bt.unmatched_ids), "no_depth": ", ".join(bt.no_depth_ids)},
        "tolerance": cell(tolerance),
        "digs_given": inputs.digs is not None,
    }


def _with_outcome(log: pd.DataFrame, status: pd.DataFrame, id_col: str, columns: list[str]) -> pd.DataFrame:
    """``log`` with ``applied_status`` and ``reason`` from the run's status file by id; a row the run did
    not record (saved after it) says so in ``applied_status``."""
    outcome = status.drop_duplicates(id_col).set_index(id_col)
    log = log.reset_index(drop=True)
    for col in ("applied_status", "reason"):
        log[col] = log[id_col].map(outcome[col])
    log["applied_status"] = log.applied_status.where(log[id_col].isin(outcome.index), UNRECORDED)
    return log[columns]


def _review_log(inputs: ReportInputs) -> pd.DataFrame:
    """Every row of ``reviews.csv`` with the outcome match recorded for it (``review_status.csv``)."""
    return _with_outcome(inputs.reviews, inputs.review_status, "decision_id", REVIEW_LOG_COLUMNS)


def _flag_log(inputs: ReportInputs) -> pd.DataFrame:
    """Every row of ``feature_flags.csv`` with the outcome classification recorded (``flag_status.csv``)."""
    return _with_outcome(inputs.flags, inputs.flag_status, "flag_id", FLAG_LOG_COLUMNS)


def pending_text(n_pending: int, n_unrecorded: int) -> str:
    """What still keeps a report from being final, in the words both the report and the CLI use: the
    matches pending review, plus (when there are any) the log rows saved after the last run, which no
    result reflects yet. The two are counted apart: a log row is not a match."""
    text = f"{n_pending} {'match' if n_pending == 1 else 'matches'} pending review"
    if n_unrecorded:
        text += f", plus {n_unrecorded} log {'row' if n_unrecorded == 1 else 'rows'} saved after the last run"
    return text


def _n_unrecorded(inputs: ReportInputs) -> int:
    """Log rows no run has recorded: saved after the last run, so not yet in any result."""
    return int((_review_log(inputs).applied_status == UNRECORDED).sum()
               + (_flag_log(inputs).applied_status == UNRECORDED).sum())


def _methods(inputs: ReportInputs) -> dict:
    tool = inputs.settings.tool.model_dump()
    config = config_rows(inputs.settings, inputs.settings_path)
    return {
        "config": table(config.assign(placeholder=config.placeholder.map({True: "PLACEHOLDER", False: ""})),
                        row_classes=["placeholder" if p else "" for p in config.placeholder]),
        "settings_path": str(inputs.settings_path),
        "tool": Table(["setting", "value"], [[f"tool.{k}", cell(v)] for k, v in tool.items()]),
        "criteria_version": inputs.criteria.version,
        "criteria_sha256": inputs.summaries["classification"].get("criteria_sha256"),
        "criteria_path": str(inputs.criteria_path),
        "criteria": table(criteria_rows(inputs.criteria)),
        "hashes": Table(["path", "sha256"], [[p, h] for p, h in input_hashes(inputs).items()]),
        "review_log": table(_review_log(inputs)),
        "flag_log": table(_flag_log(inputs)),
    }


def _environment() -> jinja2.Environment:
    return jinja2.Environment(loader=jinja2.FileSystemLoader(TEMPLATES), autoescape=True,
                              undefined=jinja2.StrictUndefined, trim_blocks=True, lstrip_blocks=True)


def render_report(inputs: ReportInputs, changes: CategoryChangeResult, header: ReportHeader,
                  generated_utc: datetime) -> str:
    """The report as one self-contained HTML document (spec 6.10), sections in ``SECTION_IDS`` order."""
    return render_html(report_context(inputs, changes, header, generated_utc))


def render_html(context: dict) -> str:
    """The HTML report from ``report_context``."""
    return _environment().get_template(TEMPLATE_NAME).render(context)


def report_context(inputs: ReportInputs, changes: CategoryChangeResult, header: ReportHeader,
                   generated_utc: datetime) -> dict:
    """Every section of the report, built once from the report inputs: the HTML template and the PDF
    (``report_pdf``) both render from it, so the two always show the same values."""
    dig = dig_list(inputs)
    n_largest = inputs.settings.report.largest_changes_n
    moved = changes.table[changes.table.moved_more_severe]
    context = {
        "label": DRAFT_LABEL,
        "label_css": _css_string(DRAFT_LABEL),
        "plotlyjs": Markup(plotly.offline.get_plotlyjs()),
        "category_colours": CATEGORY_COLOURS,
        "head": _header(inputs, header, generated_utc),
        "log_warnings": inputs.log_warnings,
        "summary": _summary(inputs, dig),
        "quality": _data_quality(inputs),
        "growth": _growth(inputs),
        "distribution": _distribution(inputs),
        "joints": _joints(inputs),
        "dig_list": table(dig, row_classes=[f"cat-{c}" for c in dig.governing_category],
                          print_columns=DIG_LIST_PRINT_COLUMNS),
        "largest": {"n": n_largest, "growth": table(largest_growth(inputs, n_largest), LARGEST_GROWTH_COLUMNS),
                    "moved": table(moved, MOVED_COLUMNS, print_columns=MOVED_PRINT_COLUMNS),
                    "n_compared": len(changes.table)},
        "charts": _charts(inputs, dig),
        "methods": _methods(inputs),
    }
    return context


# --- dig list workbook, manifest, CLI -----------------------------------------------------------------

OUTPUT_NAMES = ("category_changes.csv", "report.html", "report.pdf", "dig_list.xlsx", "anomaly_distribution.csv",
                "joint_ml_counts.csv", "anomalies_of_interest.csv", "imp_class.csv", ANALYSIS_TABLE_FILE,
                ANALYSIS_XLSX_FILE, "run_manifest.json")
DATE_COLUMNS = ("due_date", "Due Date", "Reinspection Date", "Projected Depth Condition Date",
                "Projected Burst Condition Date", "Projected Safe Condition Date")
GIT_TIMEOUT_S = 10


def _excel_value(value, column: str):
    """A cell value Excel can hold: None / NaN / NA become an empty cell, numpy scalars plain Python, and the
    date columns real dates."""
    if value is None or value is pd.NA or value is pd.NaT or (isinstance(value, float) and np.isnan(value)):
        return None
    if isinstance(value, np.generic):
        value = value.item()
    if column in DATE_COLUMNS and isinstance(value, str):
        return date.fromisoformat(value) if value else None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    return value


def dig_list_workbook(dig: pd.DataFrame, settings: Settings, label: str) -> Workbook:
    """The dig list as a workbook: row 1 the label, row 2 the header, data from row 3, each data row filled
    with its category's colour, panes frozen under the header and a filter over header and data."""
    wb = Workbook()
    ws = wb.active
    ws.title = settings.report.dig_list_sheet
    columns = list(dig.columns)
    ws.cell(1, 1, label).font = Font(bold=True)
    for j, name in enumerate(columns, start=1):
        ws.cell(2, j, name).font = Font(bold=True)
    for i, row in enumerate(dig.itertuples(index=False), start=3):
        category = row[columns.index("governing_category")]
        fill = PatternFill("solid", fgColor=CATEGORY_COLOURS[category].lstrip("#").upper())
        for j, (name, value) in enumerate(zip(columns, row), start=1):
            cell = ws.cell(i, j, _excel_value(value, name))
            if isinstance(cell.value, str):
                cell.data_type = "s"  # vendor text is never a formula, whatever it starts with
            cell.fill = fill
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(max(len(columns), 1))}{max(2, 2 + len(dig))}"
    return wb


def analysis_workbook(analysis: pd.DataFrame, label: str) -> Workbook:
    """The analysis table as a workbook: row 1 the label, row 2 the header, data from row 3, panes frozen under
    the header and right of the first two columns (Feature ID, Odometer), a filter over header and data."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    columns = list(analysis.columns)
    ws.cell(1, 1, label).font = Font(bold=True)
    for j, name in enumerate(columns, start=1):
        ws.cell(2, j, name).font = Font(bold=True)
        values = [len(name)] + [len(str(v)) for v in analysis.iloc[:200, j - 1] if v is not None and v == v]
        ws.column_dimensions[get_column_letter(j)].width = min(max(values) + 2, 40)
    for i, row in enumerate(analysis.itertuples(index=False), start=3):
        for j, (name, value) in enumerate(zip(columns, row), start=1):
            cell = ws.cell(i, j, _excel_value(value, name))
            if isinstance(cell.value, str):
                cell.data_type = "s"  # vendor text is never a formula, whatever it starts with
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A2:{get_column_letter(max(len(columns), 1))}{max(2, 2 + len(analysis))}"
    return wb


def write_dig_list(dig: pd.DataFrame, path: Path, settings: Settings, label: str) -> Path:
    dig_list_workbook(dig, settings, label).save(path)
    return Path(path)


def code_version() -> dict:
    """The package version and, when this runs from a git checkout, the commit and whether the tree is dirty
    (local git only; any failure gives None)."""
    def git(*args: str) -> str | None:
        try:
            done = subprocess.run(["git", "-C", str(PROJECT_ROOT), *args], capture_output=True, text=True,
                                  timeout=GIT_TIMEOUT_S, check=True)
        except (OSError, subprocess.SubprocessError):
            return None
        return done.stdout.strip()

    commit = git("rev-parse", "HEAD")
    status = git("status", "--porcelain") if commit else None
    return {"package": __version__, "git_commit": commit or None,
            "git_dirty": None if status is None else bool(status)}


def run_manifest(inputs: ReportInputs, outputs: list[str], created_utc: datetime) -> dict:
    """What produced the outputs: label, generator, code version, the hash of every input file, the tally and
    vendor-map hashes from the ingest manifest, the configuration files, and the output names."""
    return {
        "label": DRAFT_LABEL,
        "generator": f"ili.report {__version__}",
        "created_utc": created_utc.isoformat(timespec="seconds"),
        "code_version": code_version(),
        "inputs": input_hashes(inputs),
        "tallies": [{k: r.get(k) for k in ("run_id", "tally", "tally_sha256", "vendor_map", "vendor_map_sha256")}
                    for r in inputs.ingest_manifest.get("runs", [])],
        "config": {"settings": {"path": str(inputs.settings_path), "sha256": _sha256(inputs.settings_path)},
                   "criteria": {"path": str(inputs.criteria_path), "sha256": _sha256(inputs.criteria_path)}},
        "outputs": list(outputs),
    }


def _sheet_label(inputs: ReportInputs) -> str:
    warning = inputs.summaries["classification"].get("discovery_warning")
    return f"{DRAFT_LABEL} WARNING: {warning}" if warning else DRAFT_LABEL


REPLACE_FIRST = "dig_list.xlsx"  # the output most often held open (Excel locks it): replaced before the rest


def _write_all(writers: dict[str, Callable[[Path], object]], paths: dict[str, Path]) -> None:
    """Write each output to a temporary file beside it, then move them into place with ``os.replace``,
    ``REPLACE_FIRST`` first: a failure while writing, or replacing that file, leaves every existing output
    as it was. Temporary files are removed whatever happens."""
    temps: dict[str, Path] = {}
    try:
        for name, write in writers.items():
            temps[name] = paths[name].with_name(f".{name}.{os.getpid()}.tmp")
            write(temps[name])
        for name in sorted(temps, key=lambda n: n != REPLACE_FIRST):
            os.replace(temps[name], paths[name])
    finally:
        for tmp in temps.values():
            tmp.unlink(missing_ok=True)


def write_report(inputs: ReportInputs, header: ReportHeader, out_dir: Path,
                 generated_utc: datetime) -> tuple[dict[str, Path], CategoryChangeResult]:
    """Build every output, then write them (``_write_all``): nothing is written if any of them fails to build,
    and no existing output is replaced if ``dig_list.xlsx`` cannot be."""
    prev_dates = inputs.prev_std.run_date.dropna()
    if prev_dates.empty:
        raise ValueError("the previous run has no run date")
    changes = category_changes(inputs.match_table, inputs.remaining_life, inputs.classification, inputs.curr_std,
                               inputs.pipe, inputs.mop, inputs.segments, inputs.crossings, inputs.settings,
                               inputs.criteria, prev_dates.iloc[0])
    context = report_context(inputs, changes, header, generated_utc)
    html = render_html(context)
    pdf = build_pdf(context, inputs, generated_utc)
    workbook = dig_list_workbook(dig_list(inputs), inputs.settings, _sheet_label(inputs))
    distribution = anomaly_distribution_table(anomaly_distribution(inputs))
    joints = joint_ml_counts(inputs.weld_map, inputs.prev_std, inputs.curr_std)
    aoi = anomalies_of_interest(inputs)
    imp_rows = _imp_class_rows(inputs)
    analysis = build_analysis_table(inputs)
    analysis_book = analysis_workbook(analysis, _sheet_label(inputs))
    manifest = json.dumps(run_manifest(inputs, list(OUTPUT_NAMES), generated_utc), indent=2, allow_nan=False)
    paths = {name: Path(out_dir) / name for name in OUTPUT_NAMES}
    _write_all({
        "category_changes.csv": lambda p: changes.table.to_csv(p, index=False, lineterminator="\n"),
        "report.html": lambda p: p.write_text(html, encoding="utf-8"),
        "report.pdf": lambda p: p.write_bytes(pdf),
        "dig_list.xlsx": workbook.save,
        "anomaly_distribution.csv": lambda p: distribution.to_csv(p, index=False, lineterminator="\n"),
        "joint_ml_counts.csv": lambda p: joints.to_csv(p, index=False, lineterminator="\n"),
        "anomalies_of_interest.csv": lambda p: aoi.to_csv(p, index=False, lineterminator="\n"),
        "imp_class.csv": lambda p: imp_rows.to_csv(p, index=False, lineterminator="\n"),
        ANALYSIS_TABLE_FILE: lambda p: analysis.to_csv(p, index=False, lineterminator="\n"),
        ANALYSIS_XLSX_FILE: analysis_book.save,
        "run_manifest.json": lambda p: p.write_text(manifest, encoding="utf-8"),
    }, paths)
    return paths, changes


def _parser() -> argparse.ArgumentParser:
    ap = argparse.ArgumentParser(prog="python -m ili.report",
                                 description="HTML and PDF report and Excel dig list. " + DRAFT_LABEL)
    ap.add_argument("--prev", type=Path, required=True, help="standardized CSV of the previous run")
    ap.add_argument("--curr", type=Path, required=True, help="standardized CSV of the current run")
    ap.add_argument("--comparison", type=Path, required=True, help="the <prev>_vs_<curr> output folder")
    ap.add_argument("--pipe", type=Path, required=True, help="pipe_properties.csv")
    ap.add_argument("--mop", type=Path, required=True, help="mop_profile.csv")
    ap.add_argument("--segments", type=Path, required=True, help="hca_segments.csv")
    ap.add_argument("--crossings", type=Path, default=None, help="foreign_crossings.csv (optional)")
    ap.add_argument("--digs", type=Path, default=None, help="digs.csv, field results for the back-test (optional)")
    ap.add_argument("--repairs", type=Path, default=None,
                    help="repairs_between_runs.csv (optional; counts previously repaired anomalies)")
    ap.add_argument("--ingest-dir", type=Path, default=None,
                    help="folder with ingest_manifest.json and validation_issues.csv (default: the parent of "
                         "--curr's folder)")
    ap.add_argument("--settings", type=Path, default=DEFAULT_SETTINGS, help="settings YAML")
    ap.add_argument("--criteria", type=Path, default=None,
                    help="repair criteria YAML (default: classification.criteria_file of the settings)")
    ap.add_argument("--line-name", default=None, help="pipeline name for the report header")
    ap.add_argument("--prev-tool", default=None, help="previous run's tool type for the report header")
    ap.add_argument("--curr-tool", default=None, help="current run's tool type for the report header")
    ap.add_argument("--prepared-for", default=None, help="client (operator) named on the PDF cover")
    ap.add_argument("--prepared-by", default=None, help="author (company or engineer) named on the PDF cover")
    return ap


def main(argv: list[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    try:
        settings = load_settings(args.settings)
        ingest_dir = args.ingest_dir or args.curr.resolve().parent.parent
        inputs = load_report_inputs(args.comparison, args.prev, args.curr, ingest_dir, args.pipe, args.mop,
                                    args.segments, args.crossings, args.digs, args.settings,
                                    criteria_path(args.criteria, settings), repairs=args.repairs)
        header = ReportHeader(args.line_name, args.prev_tool, args.curr_tool, args.prepared_for, args.prepared_by)
        paths, changes = write_report(inputs, header, args.comparison, datetime.now(timezone.utc))
    except (ValueError, OSError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    summary = inputs.summaries["classification"]
    n_pending, n_unrecorded = len(pending_reviews(inputs.match_table)), _n_unrecorded(inputs)
    moved = changes.table[changes.table.moved_more_severe]
    print(DRAFT_LABEL)
    if summary.get("discovery_warning"):
        print("!!! WARNING: " + summary["discovery_warning"])
    print(f"{summary['prev_run_id']} vs {summary['curr_run_id']}: " + ", ".join(
        f"{k} {v}" for k, v in summary["counts_by_category"].items()))
    print(f"dig list: {len(dig_list(inputs))} rows")
    print(f"pending: {pending_text(n_pending, n_unrecorded)}"
          + (" (the report cannot be final)" if n_pending + n_unrecorded else ""))
    print(f"moved to a more severe category: {len(moved)} of {len(changes.table)} compared")
    for path in paths.values():
        print(f"wrote {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

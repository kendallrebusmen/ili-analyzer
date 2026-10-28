"""Load vendor ILI tallies into the standard schema (spec sections 5.1 and 6.2).

Each vendor layout is described by a YAML vendor map in ``config/vendor_maps/``:
column names, units, clock format, and the vendor's vocabulary. Ingestion
converts units, parses clock positions, maps vocabulary, and validates. Rows
that fail go to ``validation_issues.csv`` with a reason; none are dropped
silently.

DRAFT – for review by qualified personnel.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Literal

import numpy as np
import openpyxl
import pandas as pd
import yaml
from pydantic import BaseModel, ConfigDict, ValidationInfo, field_validator

import ili
from ili.config import PROJECT_ROOT
from ili.schema import STANDARD_COLUMNS, FeatureType, MlOrigin, Surface
from ili.units import CLOCK_MINUTES, FT_PER_M, IN_PER_MM, hhmm_to_clock_min

# ml_origin is derived from the feature-type code (``VendorMap.ml_origins``), not read from a column.
MAPPABLE_FIELDS = [c for c in STANDARD_COLUMNS if c not in ("run_id", "run_date", "ml_origin")]
ISSUE_COLUMNS = ["run_id", "source_row", "feature_id", "field", "value", "reason"]

# Standard field -> which vendor-map unit governs it.
UNIT_GROUP = {
    "log_dist_ft": "distance",
    "dist_from_us_gw_ft": "distance",
    "joint_length_ft": "distance",
    "wt_nom_in": "wall",
    "dent_depth_in": "dent_depth",
    "length_in": "size",
    "width_in": "size",
    "elev_ft": "elevation",
}
TO_US = {"ft": 1.0, "in": 1.0, "m": FT_PER_M, "mm": IN_PER_MM}
PLAIN_NUMBERS = ["depth_pct", "lat", "lon"]
TEXT_FIELDS = ["feature_id", "us_gw_id", "vendor_cluster_id", "comments"]
FLAG_FIELDS = ["seam_interaction", "dent_metal_loss"]


class _Strict(BaseModel):
    model_config = ConfigDict(extra="forbid")


class FileSpec(_Strict):
    format: Literal["csv", "excel"]
    sheet: str | None = None
    header_row: int = 1  # 1-based


class RunDateSpec(_Strict):
    cell: str  # e.g. "A2"
    pattern: str  # regex with one group holding an ISO date


class Units(_Strict):
    distance: Literal["ft", "m"]
    wall: Literal["in", "mm"]
    size: Literal["in", "mm"]
    dent_depth: Literal["in", "mm"]
    elevation: Literal["ft", "m"]


def _norm(token: str) -> str:
    return str(token).strip().casefold()


def _normalized_vocab(vocab: dict) -> dict:
    out: dict = {}
    for key, value in vocab.items():
        k = _norm(key)
        if k in out:
            raise ValueError(f"vocabulary key {k!r} appears twice after normalization")
        out[k] = value
    return out


class VendorMap(_Strict):
    vendor: str
    run_label: str
    file: FileSpec
    run_date: RunDateSpec | None = None
    columns: dict[str, str]
    units: Units
    clock_format: Literal["hhmm", "degrees", "decimal_hours"]
    feature_types: dict[str, FeatureType]
    surfaces: dict[str, Surface]
    ml_origins: dict[str, MlOrigin] = {}   # metal-loss feature-type code -> origin; absent code = unknown
    true_values: list[str]
    false_values: list[str]

    @field_validator("columns")
    @classmethod
    def _known_fields(cls, columns: dict[str, str]) -> dict[str, str]:
        unknown = sorted(set(columns) - set(MAPPABLE_FIELDS))
        if unknown:
            raise ValueError(f"not mappable standard fields: {', '.join(unknown)}")
        return columns

    @field_validator("feature_types", "surfaces", mode="before")
    @classmethod
    def _normalize_keys(cls, vocab: dict) -> dict:
        return _normalized_vocab(vocab)

    @field_validator("ml_origins", mode="before")
    @classmethod
    def _origin_keys_are_metal_loss(cls, vocab: dict, info: ValidationInfo) -> dict:
        types = info.data.get("feature_types")
        if types is not None:
            for key in vocab:
                if types.get(_norm(key)) is not FeatureType.metal_loss:
                    raise ValueError(f"ml_origins key {key!r} is not a metal_loss code in feature_types")
        return _normalized_vocab(vocab)

    @field_validator("true_values", "false_values")
    @classmethod
    def _normalize_flags(cls, values: list[str]) -> list[str]:
        return [_norm(v) for v in values]


def load_vendor_map(path: Path) -> VendorMap:
    with open(path, encoding="utf-8") as fh:
        return VendorMap.model_validate(yaml.safe_load(fh) or {})


def load_standardized(path: Path) -> pd.DataFrame:
    """Read a ``standardized/<run_id>.csv`` back with the schema's types: ids stay text
    (``"1"``, never ``1.0``), clock is Int64, run_date a date, flags True/False/None."""
    text_cols = ["run_id", "feature_id", "feature_type", "us_gw_id", "surface", "ml_origin", "vendor_cluster_id",
                 "comments"]
    numeric_cols = [*UNIT_GROUP, *PLAIN_NUMBERS, "clock_min", "source_row"]
    df = pd.read_csv(path, dtype={c: str for c in text_cols + FLAG_FIELDS + ["run_date"]}, keep_default_na=False,
                     na_values={c: [""] for c in numeric_cols})
    if "ml_origin" not in df.columns:  # written before M11: metal loss of unrecorded origin
        df["ml_origin"] = np.where(df.feature_type == FeatureType.metal_loss.value, MlOrigin.unknown.value, "")
    for c in text_cols:
        df[c] = df[c].astype(object).where(df[c] != "", None)
    for c in FLAG_FIELDS:
        df[c] = df[c].map({"True": True, "False": False}).astype(object).where(df[c] != "", None)
    df["run_date"] = df.run_date.map(date.fromisoformat)
    df["clock_min"] = df.clock_min.astype("Int64")
    return df


# ---------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------


def _cell(path: Path, vmap: VendorMap, ref: str) -> str | None:
    col_letters, row = re.fullmatch(r"([A-Za-z]+)(\d+)", ref).groups()
    if vmap.file.format == "excel":
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[vmap.file.sheet] if vmap.file.sheet else wb.worksheets[0]
            value = ws[ref].value
        finally:
            wb.close()
        return None if value is None else str(value)
    col = openpyxl.utils.column_index_from_string(col_letters) - 1
    with open(path, encoding="utf-8-sig", newline="") as fh:
        for i, line in enumerate(csv.reader(fh), start=1):
            if i == int(row):
                return line[col] if col < len(line) else None
    return None


def _excel_text(cell) -> str:
    """A cell as the text a person sees in Excel: percent cells as percent, times as h:mm."""
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(sep=" ")
    if isinstance(v, dt.time):
        minutes = v.hour * 60 + v.minute + round(v.second / 60)
        return f"{minutes // 60}:{minutes % 60:02d}"
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        if "%" in (cell.number_format or ""):
            return f"{v * 100:.12g}"
        return str(v)
    return str(v)


def _excel_records(path: Path, vmap: VendorMap) -> list[tuple[int, list[str]]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        if vmap.file.sheet and vmap.file.sheet not in wb.sheetnames:
            raise ValueError(f"{path.name}: sheet {vmap.file.sheet!r} not found (has {wb.sheetnames})")
        ws = wb[vmap.file.sheet] if vmap.file.sheet else wb.worksheets[0]
        return [(row[0].row, [_excel_text(c) for c in row]) for row in ws.iter_rows(min_row=vmap.file.header_row)]
    finally:
        wb.close()


def _csv_records(path: Path, vmap: VendorMap) -> list[tuple[int, list[str]]]:
    """Records with the physical line each starts on (blank lines and multi-line cells counted)."""
    records = []
    with open(path, encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        last_line = 0
        for fields in reader:
            start, last_line = last_line + 1, reader.line_num
            if start >= vmap.file.header_row:
                records.append((start, fields))
    if not records or records[0][0] != vmap.file.header_row:
        raise ValueError(f"{path.name}: no header record starts on line {vmap.file.header_row}")
    return records


def read_tally(path: Path, vmap: VendorMap) -> tuple[pd.DataFrame, date | None]:
    """Raw tally as strings (blank = ""), with ``source_row`` = the 1-based physical row the
    record starts on, and the run date if the vendor map says where to find it.

    Rows with no content at all are skipped (there is nothing in them to lose).
    """
    path = Path(path)
    records = _csv_records(path, vmap) if vmap.file.format == "csv" else _excel_records(path, vmap)
    header = [h.strip() for h in records[0][1]]
    rows, source_rows = [], []
    for line, fields in records[1:]:
        fields = [f.strip() for f in fields]
        if not any(fields):
            continue
        if len(fields) > len(header):
            if any(fields[len(header):]):
                raise ValueError(f"{path.name}: row {line} has more values than the header has columns")
            fields = fields[: len(header)]
        rows.append(fields + [""] * (len(header) - len(fields)))
        source_rows.append(line)
    missing = [v for v in vmap.columns.values() if v not in header]
    if missing:
        raise ValueError(f"{path.name}: columns required by vendor map {vmap.vendor!r} are missing: "
                         + ", ".join(repr(m) for m in missing))
    repeated = sorted({v for v in vmap.columns.values() if header.count(v) > 1})
    if repeated:
        raise ValueError(f"{path.name}: mapped column(s) appear more than once: " + ", ".join(map(repr, repeated)))
    raw = pd.DataFrame(rows, columns=header, dtype=object)
    raw = raw.loc[:, ~raw.columns.duplicated()]
    raw["source_row"] = np.array(source_rows, dtype=int)

    run_date = None
    if vmap.run_date is not None:
        text = _cell(path, vmap, vmap.run_date.cell) or ""
        match = re.search(vmap.run_date.pattern, text)
        if not match:
            raise ValueError(f"{path.name}: run date not found in cell {vmap.run_date.cell} ({text!r})")
        run_date = date.fromisoformat(match.group(1))
    return raw, run_date


# ---------------------------------------------------------------------------
# Standardizing
# ---------------------------------------------------------------------------


def parse_clock(values: pd.Series, fmt: str) -> tuple[pd.Series, pd.Series]:
    """Clock positions to minutes from 12:00. Returns (Int64 minutes, unparseable mask).

    Valid numeric ranges are 0-360 degrees and 0-<13 decimal hours. A full turn (360 deg,
    12.x h on a 12-hour face, like hh:mm's 12:xx) wraps into 0-719. Out-of-range numbers
    are kept unwrapped so validation reports them rather than wrapping them silently.
    """
    values = values.astype(object).fillna("").astype(str).str.strip()
    blank = values == ""
    minutes = pd.Series(pd.NA, index=values.index, dtype="Int64")
    bad = pd.Series(False, index=values.index)
    if fmt == "hhmm":
        for idx, text in values[~blank].items():
            try:
                minutes[idx] = hhmm_to_clock_min(text)
            except ValueError:
                bad[idx] = True
        return minutes, bad
    per_unit = {"degrees": 2.0, "decimal_hours": 60.0}[fmt]
    num = pd.to_numeric(values.where(~blank), errors="coerce")
    bad = ~blank & num.isna()
    in_range = (num >= 0) & ((num <= 360) if fmt == "degrees" else (num < 13))
    m = np.rint(num * per_unit)
    m = m.where(~in_range, m % CLOCK_MINUTES)
    ok = m.notna()
    minutes[ok] = m[ok].astype(int)
    return minutes, bad


def _issues(run_id, rows: pd.DataFrame, mask, field: str, values, reason: str) -> pd.DataFrame:
    mask = np.asarray(mask, dtype=bool)
    return pd.DataFrame({
        "run_id": run_id,
        "source_row": rows.source_row.to_numpy()[mask],
        "feature_id": rows.feature_id.to_numpy()[mask],
        "field": field,
        "value": np.asarray(values, dtype=object)[mask],
        "reason": reason,
    }, columns=ISSUE_COLUMNS)


def _vocab(values: pd.Series, mapping: dict) -> tuple[list, np.ndarray]:
    keys = [str(v).strip().casefold() for v in values]
    mapped = [None if k == "" else (mapping[k].value if k in mapping else None) for k in keys]
    unknown = np.array([k != "" and k not in mapping for k in keys])
    return mapped, unknown


def standardize(raw: pd.DataFrame, vmap: VendorMap, run_id: str, run_date: date) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Rename, convert units, parse clocks and flags, and map vocabulary.

    Returns the standard frame (``STANDARD_COLUMNS`` + ``source_row``, one row per input
    row) and the issues found while parsing. Rows with parse issues are kept here, with the
    bad field blank; ``ingest_tally`` excludes them.
    """
    src = {std: raw[col] if col else pd.Series("", index=raw.index)
           for std, col in ((f, vmap.columns.get(f)) for f in MAPPABLE_FIELDS)}
    out = pd.DataFrame(index=raw.index)
    out["run_id"] = run_id
    out["run_date"] = run_date
    issues = []
    ident = pd.DataFrame({"source_row": raw.source_row,
                          "feature_id": src["feature_id"].replace("", None)})

    for f in TEXT_FIELDS:
        out[f] = src[f].astype(object).where(src[f] != "", None)

    out["feature_type"], unknown = _vocab(src["feature_type"], vmap.feature_types)
    issues.append(_issues(run_id, ident, unknown, "feature_type", src["feature_type"], "unknown_feature_type"))
    out["surface"], unknown = _vocab(src["surface"], vmap.surfaces)
    issues.append(_issues(run_id, ident, unknown, "surface", src["surface"], "unknown_surface"))
    codes = src["feature_type"].astype(str).str.strip().str.casefold()
    is_ml = (out["feature_type"] == FeatureType.metal_loss.value).to_numpy()
    out["ml_origin"] = pd.Series([vmap.ml_origins.get(k, MlOrigin.unknown).value if ml else None
                                  for k, ml in zip(codes, is_ml)], index=raw.index, dtype=object)

    for f in [*UNIT_GROUP, *PLAIN_NUMBERS]:
        text = src[f]
        num = pd.to_numeric(text.where(text != ""), errors="coerce")
        issues.append(_issues(run_id, ident, (text != "") & num.isna(), f, text, "unparseable_number"))
        if f in UNIT_GROUP:
            num = num * TO_US[getattr(vmap.units, UNIT_GROUP[f])]
        out[f] = num.astype(float)

    out["clock_min"], bad = parse_clock(src["clock_min"], vmap.clock_format)
    issues.append(_issues(run_id, ident, bad, "clock_min", src["clock_min"], "unparseable_clock"))

    for f in FLAG_FIELDS:
        keys = src[f].astype(str).str.strip().str.casefold()
        is_true, is_false = keys.isin(vmap.true_values), keys.isin(vmap.false_values)
        out[f] = [True if t else (False if fl else None) for t, fl in zip(is_true, is_false)]
        out[f] = out[f].astype(object)
        issues.append(_issues(run_id, ident, (keys != "") & ~is_true & ~is_false, f, src[f], "unparseable_flag"))

    out["source_row"] = raw.source_row.to_numpy()
    issues_df = pd.concat([i for i in issues if len(i)], ignore_index=True) if any(len(i) for i in issues) \
        else pd.DataFrame(columns=ISSUE_COLUMNS)
    return out[STANDARD_COLUMNS + ["source_row"]].reset_index(drop=True), issues_df


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------

REQUIRED_FIELDS = ["feature_id", "feature_type", "log_dist_ft"]

# field -> (lower bound, upper bound, lower bound exclusive). Non-finite values always fail.
# Matches the StandardFeature constraints plus physical limits; log distance and elevation
# may be negative (positions upstream of the odometer zero, below datum).
VALUE_BOUNDS = {
    "log_dist_ft": (None, None, False),
    "dist_from_us_gw_ft": (0.0, None, False),
    "joint_length_ft": (0.0, None, True),
    "wt_nom_in": (0.0, None, True),
    "dent_depth_in": (0.0, None, False),
    "length_in": (0.0, None, False),
    "width_in": (0.0, None, False),
    "lat": (-90.0, 90.0, False),
    "lon": (-180.0, 180.0, False),
    "elev_ft": (None, None, False),
}


def _longest_increasing(values: np.ndarray) -> np.ndarray:
    """Mask of one longest strictly increasing subsequence (O(n log n) patience sort)."""
    n = len(values)
    tails: list[int] = []  # index of the smallest tail of each length
    prev = np.full(n, -1)
    for i, v in enumerate(values):
        lo, hi = 0, len(tails)
        while lo < hi:
            mid = (lo + hi) // 2
            if values[tails[mid]] < v:
                lo = mid + 1
            else:
                hi = mid
        prev[i] = tails[lo - 1] if lo > 0 else -1
        if lo == len(tails):
            tails.append(i)
        else:
            tails[lo] = i
    keep = np.zeros(n, dtype=bool)
    i = tails[-1] if tails else -1
    while i >= 0:
        keep[i] = True
        i = prev[i]
    return keep


def validate(std: pd.DataFrame, run_id: str, parse_issues: pd.DataFrame | None = None
             ) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Split a standardized run into clean rows and issue rows (spec 5.1).

    Each problem is one issue line; a row with any problem, including a parse issue from
    ``standardize``, is excluded from the clean output exactly once.
    """
    parse_issues = parse_issues if parse_issues is not None else pd.DataFrame(columns=ISSUE_COLUMNS)
    already = set(zip(parse_issues.source_row, parse_issues.field))
    issues = [parse_issues]

    for f in REQUIRED_FIELDS:
        missing = std[f].isna().to_numpy() & np.array([(r, f) not in already for r in std.source_row], dtype=bool)
        issues.append(_issues(run_id, std, missing, f, std[f], "missing_required"))

    depth = std.depth_pct
    issues.append(_issues(run_id, std, depth.notna() & ~depth.between(0, 100), "depth_pct", depth,
                          "depth_out_of_range"))
    clock = std.clock_min.astype("Float64")
    bad_clock = (clock.notna() & ((clock < 0) | (clock >= CLOCK_MINUTES))).fillna(False)
    issues.append(_issues(run_id, std, bad_clock, "clock_min", std.clock_min.astype(object), "clock_out_of_range"))

    dup = std.feature_id.notna() & std.feature_id.duplicated(keep="first")
    issues.append(_issues(run_id, std, dup, "feature_id", std.feature_id, "duplicate_feature_id"))

    for f, (lo, hi, lo_open) in VALUE_BOUNDS.items():
        v = std[f]
        out = v.notna() & ~np.isfinite(v.fillna(0))
        if lo is not None:
            out |= (v <= lo) if lo_open else (v < lo)
        if hi is not None:
            out |= v > hi
        issues.append(_issues(run_id, std, out.fillna(False), f, v, "value_out_of_range"))

    # Girth welds (file order) must increase in log distance. Keep the longest strictly
    # increasing sequence and exclude every other weld, so the clean set is always monotonic
    # whatever went wrong (swap, displaced block, odometer reset). Alignment treats the
    # excluded welds as missed welds.
    welds = std[(std.feature_type == "girth_weld") & std.log_dist_ft.notna()].sort_values("source_row")
    keep = _longest_increasing(welds.log_dist_ft.to_numpy())
    bad_rows = set(welds.source_row[~keep])
    not_mono = std.source_row.isin(bad_rows).to_numpy()
    issues.append(_issues(run_id, std, not_mono, "log_dist_ft", std.log_dist_ft, "girth_weld_not_monotonic"))

    found = [i for i in issues if len(i)]
    all_issues = (pd.concat(found, ignore_index=True) if found else pd.DataFrame(columns=ISSUE_COLUMNS))
    all_issues = all_issues.sort_values("source_row", kind="stable").reset_index(drop=True)
    clean = std[~std.source_row.isin(set(all_issues.source_row))].reset_index(drop=True)
    return clean, all_issues


# ---------------------------------------------------------------------------
# Orchestration, output, CLI
# ---------------------------------------------------------------------------

DRAFT_LABEL = "DRAFT – for review by qualified personnel."


@dataclass
class IngestResult:
    run_id: str
    run_date: date
    standard: pd.DataFrame  # clean rows: STANDARD_COLUMNS + source_row
    issues: pd.DataFrame  # ISSUE_COLUMNS
    tally_path: Path
    vendor_map_path: Path
    rows_read: int


def ingest_tally(path: Path, vendor_map_path: Path, run_date: date | None = None) -> IngestResult:
    """Read, standardize, and validate one tally. The run date comes from ``run_date``,
    else from the file (if the vendor map locates it); if both exist they must agree."""
    vmap = load_vendor_map(vendor_map_path)
    raw, file_date = read_tally(path, vmap)
    if run_date and file_date and run_date != file_date:
        raise ValueError(f"{Path(path).name}: run_date {run_date} conflicts with {file_date} in the file")
    run_date = run_date or file_date
    if run_date is None:
        raise ValueError(f"{Path(path).name}: run_date is not in the file; supply it")
    run_id = f"{run_date.year}_{vmap.run_label}"
    std, parse_issues = standardize(raw, vmap, run_id, run_date)
    clean, issues = validate(std, run_id, parse_issues)
    return IngestResult(run_id, run_date, clean, issues, Path(path), Path(vendor_map_path), len(raw))


def _sha256(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_ingest(results: list[IngestResult], out_dir: Path = PROJECT_ROOT / "outputs") -> dict[str, Path]:
    ids = [r.run_id for r in results]
    dupes = sorted({i for i in ids if ids.count(i) > 1})
    if dupes:
        raise ValueError(f"duplicate run_id: {', '.join(dupes)}")
    out_dir = Path(out_dir)
    (out_dir / "standardized").mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}
    for r in results:
        rel = f"standardized/{r.run_id}.csv"
        r.standard.to_csv(out_dir / rel, index=False, lineterminator="\n")
        paths[rel] = out_dir / rel
    # Issues and manifest entries of runs ingested by earlier invocations are kept; this
    # invocation's runs replace their own earlier entries.
    issues_path = out_dir / "validation_issues.csv"
    manifest_path = out_dir / "ingest_manifest.json"
    earlier = [pd.DataFrame(columns=ISSUE_COLUMNS)]
    if issues_path.exists():
        old = pd.read_csv(issues_path, dtype=str, keep_default_na=False)
        earlier.append(old[~old.run_id.isin(ids)].astype({"source_row": int}))
    issues = [*earlier, *(r.issues for r in results)]
    combined = pd.concat([i for i in issues if len(i)] or [issues[0]], ignore_index=True)
    combined.to_csv(issues_path, index=False, lineterminator="\n")
    paths["validation_issues.csv"] = issues_path
    earlier_runs = []
    if manifest_path.exists():
        earlier_runs = [r for r in json.loads(manifest_path.read_text(encoding="utf-8")).get("runs", [])
                        if r.get("run_id") not in ids]
    manifest = {
        "label": DRAFT_LABEL,
        "generator": f"ili.ingest {ili.__version__}",
        "created_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "runs": earlier_runs + [{
            "run_id": r.run_id,
            "run_date": r.run_date.isoformat(),
            "tally": str(r.tally_path),
            "tally_sha256": _sha256(r.tally_path),
            "vendor_map": str(r.vendor_map_path),
            "vendor_map_sha256": _sha256(r.vendor_map_path),
            "rows_read": r.rows_read,
            "rows_clean": len(r.standard),
            "issues": len(r.issues),
        } for r in results],
    }
    (out_dir / "ingest_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    paths["ingest_manifest.json"] = out_dir / "ingest_manifest.json"
    return paths


def _parser() -> argparse.ArgumentParser:
    ap = argparse.ArgumentParser(prog="python -m ili.ingest",
                                 description="Load ILI tallies into the standard schema. " + DRAFT_LABEL)
    ap.add_argument("--run", action="append", nargs="+", required=True, metavar="ARG",
                    help="TALLY VENDOR_MAP [RUN_DATE as YYYY-MM-DD]; repeat for each run")
    ap.add_argument("--out", type=Path, default=PROJECT_ROOT / "outputs", help="output directory")
    return ap


def main(argv: list[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    try:
        results = []
        for spec in args.run:
            if len(spec) not in (2, 3):
                raise ValueError(f"--run takes TALLY VENDOR_MAP [RUN_DATE], got {len(spec)} values")
            run_date = date.fromisoformat(spec[2]) if len(spec) == 3 else None
            results.append(ingest_tally(Path(spec[0]), Path(spec[1]), run_date))
        write_ingest(results, args.out)
    except ValueError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    print(DRAFT_LABEL)
    for r in results:
        print(f"{r.run_id}: {r.rows_read} rows read, {len(r.standard)} clean, {len(r.issues)} issues")
    print(f"wrote {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

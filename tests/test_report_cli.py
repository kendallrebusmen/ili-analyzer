import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from ili import __version__, classify, config, remaining_life
from ili.align_welds import DRAFT_LABEL
from ili.report import code_version, main, write_dig_list
from ili.report_charts import CATEGORY_COLOURS
from ili.report_data import DIG_LIST_COLUMNS
from test_classify_cli import _classify_args
from test_rl_cli import _args as rl_args
from test_rl_cli import _prepare


@pytest.fixture(scope="module")
def built(synth, tmp_path_factory):
    _, data = synth
    out = tmp_path_factory.mktemp("report_cli") / "out"
    std, folder = _prepare(data, out)
    assert remaining_life.main(rl_args(std, folder, data, out)) == 0
    assert classify.main(_classify_args(std, folder, data, out)) == 0
    return data, out, std, folder


def _argv(built, *extra, mop="mop_profile.csv"):
    data, out, std, folder = built
    ref = data / "reference"
    return ["--prev", str(std / "2021_vendorA.csv"), "--curr", str(std / "2026_vendorB.csv"),
            "--comparison", str(folder), "--pipe", str(ref / "pipe_properties.csv"), "--mop", str(ref / mop),
            "--segments", str(ref / "hca_segments.csv"), "--crossings", str(ref / "foreign_crossings.csv"),
            "--digs", str(ref / "digs.csv"), *extra]


def test_cli_writes_outputs(built, capsys):
    _, _, _, folder = built
    capsys.readouterr()
    assert main(_argv(built, "--line-name", "Line 7")) == 0
    printed = capsys.readouterr().out
    names = ["category_changes.csv", "report.html", "dig_list.xlsx", "anomaly_distribution.csv",
             "joint_ml_counts.csv", "anomalies_of_interest.csv", "run_manifest.json"]
    assert all((folder / n).is_file() for n in names)
    assert DRAFT_LABEL in printed and "pending" in printed and "dig list" in printed

    cls = pd.read_csv(folder / "classification.csv")
    dig = cls[cls.governing_category != "evaluate"]
    ws = openpyxl.load_workbook(folder / "dig_list.xlsx").active
    assert ws.title == config.load_settings().report.dig_list_sheet
    assert "DRAFT" in ws["A1"].value
    assert ws.freeze_panes == "A3"
    assert [c.value for c in ws[2]] == list(DIG_LIST_COLUMNS)
    assert ws.max_row - 2 == len(dig) > 0
    assert ws.auto_filter.ref == f"A2:{openpyxl.utils.get_column_letter(len(DIG_LIST_COLUMNS))}{ws.max_row}"
    category = ws.cell(3, 1).value
    assert ws.cell(3, 1).fill.fgColor.rgb[-6:] == CATEGORY_COLOURS[category].lstrip("#").upper()

    manifest = json.loads((folder / "run_manifest.json").read_text(encoding="utf-8"))
    assert DRAFT_LABEL == manifest["label"] and manifest["generator"] == f"ili.report {__version__}"
    assert manifest["inputs"]
    for path, digest in manifest["inputs"].items():
        assert hashlib.sha256(open(path, "rb").read()).hexdigest() == digest
    assert set(names) <= set(manifest["outputs"])
    assert set(manifest["code_version"]) == {"package", "git_commit", "git_dirty"}
    assert manifest["config"]["settings"]["sha256"] and manifest["config"]["criteria"]["sha256"]
    assert {r["run_id"] for r in manifest["tallies"]} == {"2021_vendorA", "2026_vendorB"}
    assert "Line 7" in (folder / "report.html").read_text(encoding="utf-8")


def test_report_cli_writes_all_outputs(built, capsys):
    from ili.report import OUTPUT_NAMES
    _, _, _, folder = built
    assert main(_argv(built, "--line-name", "Line 7")) == 0
    capsys.readouterr()
    assert len(OUTPUT_NAMES) == 7
    manifest = json.loads((folder / "run_manifest.json").read_text(encoding="utf-8"))
    for name in OUTPUT_NAMES:
        assert (folder / name).is_file()
        assert name in manifest["outputs"]


def test_empty_dig_list_workbook(tmp_path):
    path = tmp_path / "dig.xlsx"
    write_dig_list(pd.DataFrame(columns=list(DIG_LIST_COLUMNS)), path, config.load_settings(), DRAFT_LABEL)
    ws = openpyxl.load_workbook(path).active
    assert [c.value for c in ws[2]] == list(DIG_LIST_COLUMNS)
    assert ws.max_row == 2 and ws.freeze_panes == "A3"
    assert ws.auto_filter.ref == f"A2:{openpyxl.utils.get_column_letter(len(DIG_LIST_COLUMNS))}2"


def test_dig_list_dates_and_blanks(tmp_path):
    row = {c: None for c in DIG_LIST_COLUMNS}
    row.update(governing_category="60_day", match_id="M1", due_date="2027-01-02", p_fail=float("nan"))
    path = tmp_path / "dig.xlsx"
    write_dig_list(pd.DataFrame([row]), path, config.load_settings(), DRAFT_LABEL)
    ws = openpyxl.load_workbook(path).active
    col = {c.value: c.column for c in ws[2]}
    assert isinstance(ws.cell(3, col["due_date"]).value, datetime)
    assert ws.cell(3, col["p_fail"]).value is None and ws.cell(3, col["rl_yr"]).value is None


def test_code_version_shape():
    v = code_version()
    assert v["package"] == __version__
    assert v["git_commit"] is None or len(v["git_commit"]) == 40
    assert v["git_dirty"] is None or isinstance(v["git_dirty"], bool)


def test_cli_missing_input_exit_2(built, capsys):
    capsys.readouterr()
    assert main(_argv(built, mop="nope.csv")) == 2
    assert "error:" in capsys.readouterr().err


def test_cli_stale_criteria_exit_2(built, tmp_path, capsys):
    stale = tmp_path / "criteria.yaml"
    stale.write_text(config.DEFAULT_CRITERIA.read_text(encoding="utf-8") + "\n# edited\n", encoding="utf-8")
    capsys.readouterr()
    assert main(_argv(built, "--criteria", str(stale))) == 2
    err = capsys.readouterr().err
    assert "error:" in err and "changed since" in err


def test_dig_list_strings_are_never_formulas(tmp_path):
    texts = ["=1+1", "+1", "-1", "@SUM(1)"]
    rows = []
    for t in texts:
        row = {c: None for c in DIG_LIST_COLUMNS}
        row.update(governing_category="immediate", curr_feature_id=t, review_reasons=t)
        rows.append(row)
    path = tmp_path / "dig.xlsx"
    write_dig_list(pd.DataFrame(rows), path, config.load_settings(), DRAFT_LABEL)
    ws = openpyxl.load_workbook(path, data_only=False).active
    col = {c.value: c.column for c in ws[2]}
    for i, t in enumerate(texts, start=3):
        for name in ("curr_feature_id", "review_reasons"):
            cell = ws.cell(i, col[name])
            assert cell.data_type == "s" and cell.value == t


def test_dig_list_nat_is_empty_cell(tmp_path):
    row = {c: None for c in DIG_LIST_COLUMNS}
    row.update(governing_category="60_day", due_date=pd.NaT)
    path = tmp_path / "dig.xlsx"
    write_dig_list(pd.DataFrame([row]), path, config.load_settings(), DRAFT_LABEL)
    ws = openpyxl.load_workbook(path).active
    col = {c.value: c.column for c in ws[2]}
    assert ws.cell(3, col["due_date"]).value is None


@pytest.mark.parametrize("target", ["render_report", "dig_list_workbook", "run_manifest"])
def test_no_file_written_when_a_build_step_fails(built, tmp_path, monkeypatch, target):
    from ili import report
    from ili.report import ReportHeader, write_report
    from ili.report_data import load_report_inputs
    data, out, std, folder = built
    ref = data / "reference"
    inputs = load_report_inputs(folder, std / "2021_vendorA.csv", std / "2026_vendorB.csv", out,
                                ref / "pipe_properties.csv", ref / "mop_profile.csv", ref / "hca_segments.csv",
                                None, None, config.DEFAULT_SETTINGS, config.DEFAULT_CRITERIA)

    def boom(*a, **k):
        raise ValueError("boom")
    monkeypatch.setattr(report, target, boom)
    dest = tmp_path / "empty"
    dest.mkdir()
    with pytest.raises(ValueError, match="boom"):
        write_report(inputs, ReportHeader(None, None, None), dest, datetime.now(timezone.utc))
    assert list(dest.iterdir()) == []


def _inputs(built):
    from ili.report_data import load_report_inputs
    data, out, std, folder = built
    ref = data / "reference"
    return load_report_inputs(folder, std / "2021_vendorA.csv", std / "2026_vendorB.csv", out,
                              ref / "pipe_properties.csv", ref / "mop_profile.csv", ref / "hca_segments.csv",
                              None, None, config.DEFAULT_SETTINGS, config.DEFAULT_CRITERIA)


def test_category_changes_csv_has_lf_line_endings(built, tmp_path):
    from ili.report import ReportHeader, write_report
    paths, _ = write_report(_inputs(built), ReportHeader(None, None, None), tmp_path, datetime.now(timezone.utc))
    data = paths["category_changes.csv"].read_bytes()
    assert data.count(b"\n") > 1 and b"\r" not in data


def test_locked_dig_list_leaves_every_output_untouched(built, tmp_path, monkeypatch):
    """Excel holding dig_list.xlsx open makes writing or replacing it fail: no other output may have been
    replaced."""
    import os

    from ili.report import OUTPUT_NAMES, ReportHeader, write_report
    for name in OUTPUT_NAMES:
        (tmp_path / name).write_bytes(b"old " + name.encode())
    real_replace, real_save = os.replace, openpyxl.Workbook.save

    def check(path):
        if Path(path).name == "dig_list.xlsx":
            raise PermissionError(13, "The process cannot access the file", str(path))

    def replace(src, dst):
        check(dst)
        return real_replace(src, dst)

    def save(wb, path):
        check(path)
        return real_save(wb, path)
    monkeypatch.setattr(os, "replace", replace)
    monkeypatch.setattr(openpyxl.Workbook, "save", save)
    with pytest.raises(PermissionError):
        write_report(_inputs(built), ReportHeader(None, None, None), tmp_path, datetime.now(timezone.utc))
    assert sorted(p.name for p in tmp_path.iterdir()) == sorted(OUTPUT_NAMES)  # no temp file left behind
    for name in OUTPUT_NAMES:
        assert (tmp_path / name).read_bytes() == b"old " + name.encode(), name


@pytest.mark.parametrize("warning", [None, "Discovery date not entered."])
def test_dig_list_row_1_carries_the_discovery_warning(built, tmp_path, warning):
    import dataclasses

    from ili.report import ReportHeader, write_report
    inputs = _inputs(built)
    cls = {**inputs.summaries["classification"], "discovery_warning": warning}
    inputs = dataclasses.replace(inputs, summaries={**inputs.summaries, "classification": cls})
    paths, _ = write_report(inputs, ReportHeader(None, None, None), tmp_path, datetime.now(timezone.utc))
    a1 = openpyxl.load_workbook(paths["dig_list.xlsx"]).active["A1"].value
    assert a1 == (DRAFT_LABEL if warning is None else f"{DRAFT_LABEL} WARNING: {warning}")


def test_report_writes_anomaly_distribution(built):
    _, _, _, folder = built
    assert main(_argv(built)) == 0
    out = pd.read_csv(folder / "anomaly_distribution.csv", dtype={"bin": str})
    settings = config.load_settings()
    n_depth, n_dent = len(settings.report.depth_bins_pct_wt), len(settings.report.dent_bins_pct_od)
    assert list(out.columns[:2]) == ["table", "run_id"]
    assert set(out.table) == {"depth", "dent"}
    assert (out.table == "depth").sum() == 2 * 3 * n_depth and (out.table == "dent").sum() == 2 * n_dent
    depth = out[(out.table == "depth") & (out.location == "all")]
    assert depth.total.sum() > 0 and {"2021_vendorA", "2026_vendorB"} == set(depth.run_id)
    assert settings.classification.girth_weld_proximity_in is None  # so the girth-weld rows are not evaluated
    gw = out[(out.table == "depth") & (out.location == "girth_weld")]
    assert (gw.total == 0).all() and (gw.n_location_unknown > 0).all()
    assert (depth.n_location_unknown == 0).all()
    manifest = json.loads((folder / "run_manifest.json").read_text(encoding="utf-8"))
    assert "anomaly_distribution.csv" in manifest["outputs"]


def test_report_writes_joint_counts(built):
    _, _, _, folder = built
    assert main(_argv(built)) == 0
    out = pd.read_csv(folder / "joint_ml_counts.csv", dtype={"prev_gw_ids": str, "curr_gw_ids": str})
    weld_map = pd.read_csv(folder / "weld_map.csv")
    assert len(out) == len(weld_map) and list(out.columns[:5]) == [
        "section_id", "match_type", "prev_gw_ids", "curr_gw_ids", "curr_start_ft"]
    assert out.curr_total.sum() > 0 and out.prev_total.sum() > 0
    both = out.dropna(subset=["curr_total"])
    assert (both.curr_total == both.curr_ext + both.curr_int + both.curr_other).all()
    assert out.prev_total[out.match_type == "inserted"].isna().all()
    manifest = json.loads((folder / "run_manifest.json").read_text(encoding="utf-8"))
    assert "joint_ml_counts.csv" in manifest["outputs"]
    html = (folder / "report.html").read_text(encoding="utf-8")
    assert "Per-joint metal-loss counts" in html


def test_report_writes_aoi(built):
    _, _, _, folder = built
    assert main(_argv(built)) == 0
    out = pd.read_csv(folder / "anomalies_of_interest.csv")
    assert list(out.columns) == ["rank", "start_ft", "end_ft", "surfaces", "origins", "max_depth_pct",
                                 "n_anomalies", "min_rl_yr", "hca", "description"]
    assert 0 < len(out) <= config.load_settings().report.aoi_count
    assert list(out["rank"]) == list(range(1, len(out) + 1))
    manifest = json.loads((folder / "run_manifest.json").read_text(encoding="utf-8"))
    assert "anomalies_of_interest.csv" in manifest["outputs"]
    assert "Anomalies of interest" in (folder / "report.html").read_text(encoding="utf-8")

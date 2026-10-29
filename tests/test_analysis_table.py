"""The analysis table: the consultant's classified-anomalies layout, built from what the stages wrote."""

import dataclasses
from datetime import datetime, timezone

import numpy as np
import openpyxl
import pandas as pd
import pytest

from ili.align_welds import DRAFT_LABEL
from ili.analysis_table import (ANALYSIS_COLUMNS, ANALYSIS_TABLE_FILE, ANALYSIS_XLSX_FILE, build_analysis_table,
                                load_analysis_table)
from test_report_inputs import _load, built  # noqa: F401  (built is the module fixture)

NAMES = [n for n, _ in ANALYSIS_COLUMNS]


@pytest.fixture(scope="module")
def inputs(built):  # noqa: F811
    return _load(built)


@pytest.fixture(scope="module")
def table(inputs):
    return build_analysis_table(inputs)


def test_columns_in_the_consultants_order(table):
    assert list(table.columns) == NAMES
    assert {g for _, g in ANALYSIS_COLUMNS} == {"Identity & location", "Growth", "Dents", "Size & pipe", "Strength",
                                                "Projection", "Life & class", "Nemi"}
    for name in ("Feature ID", "Odometer (ft)", "Corrosion Growth Rate (MPY)", "Mod B31G Burst Pressure (psi)",
                 "Projected Reinspection ERF", "Remaining Life (Years)", "IMP Classification", "Requires Review"):
        assert name in NAMES


def test_one_row_per_classification_row_in_distance_order(inputs, table):
    assert len(table) == len(inputs.classification) > 0
    assert set(table["Match ID"]) == set(inputs.classification.match_id)
    assert table["Odometer (ft)"].dropna().is_monotonic_increasing


def test_values_match_their_sources(inputs, table):
    cls = inputs.classification.set_index("match_id")
    rl = inputs.remaining_life.drop_duplicates("match_id").set_index("match_id")
    growth = inputs.growth.drop_duplicates("match_id").set_index("match_id")
    mt = inputs.match_table.drop_duplicates("match_id").set_index("match_id")
    metal = table[table["Feature Category"] == "Metal loss"]
    assert len(metal) > 5
    for _, row in metal.head(40).iterrows():
        m = row["Match ID"]
        for col, src in (("ERF", rl.erf), ("Remaining Life (Years)", rl.rl_yr),
                         ("Corrosion Growth Rate (MPY)", growth.governing_mpy),
                         ("Metal Loss Depth (%)", cls.depth_pct_reported), ("Odometer (ft)", mt.curr_log_dist_ft),
                         ("Remaining Life, Factored (Years)", rl.rl_factored_yr)):
            want = src.get(m)
            if pd.isna(want):
                assert pd.isna(row[col]), (m, col)
            else:
                assert row[col] == pytest.approx(float(want)), (m, col)
        assert row["Feature ID"] == cls.curr_feature_id[m]
        assert row["Requires Review"] in ("Y", "N")
    joint = metal.dropna(subset=["Joint Length (ft)", "U/S Weld Distance (ft)"])
    assert len(joint) > 0
    assert np.allclose(joint["D/S Weld Distance (ft)"], joint["Joint Length (ft)"] - joint["U/S Weld Distance (ft)"])
    both = metal.dropna(subset=["Metal Loss Depth (%)", "Prev. Anomaly Depth (%)"])
    assert len(both) > 0
    assert np.allclose(both["Change in Anomaly Depth (%)"], both["Metal Loss Depth (%)"] - both["Prev. Anomaly Depth (%)"])
    assert set(metal["Internal / External"]) <= {"Internal", "External", "Unknown"}
    assert metal["Orientation (deg)"].dropna().between(0, 359.5).all()
    assert metal["Orientation (o'clock)"].str.match(r"^\d{1,2}:\d\d$").all()
    assert set(table["Is Dent on Top of Pipe"].dropna()) <= {"Y", "N"}
    assert (table["Depth Tolerance (%)"].dropna() >= 0).all() and table["Depth Tolerance (%)"].notna().any()
    assert table["Prev. Measurement Date"].dropna().str.match(r"^\d{4}-\d\d-\d\d$").all()


def test_dents_have_blank_growth_and_remaining_life(table):
    dents = table[table["Feature Category"] != "Metal loss"]
    if dents.empty:
        pytest.skip("the synthetic line has no dents or gouges")
    for col in ("Corrosion Growth Rate (MPY)", "Growth Rate Source", "ERF", "FPR", "Remaining Life (Years)",
                "Mod B31G Burst Pressure (psi)", "Projected Reinspection ERF"):
        assert dents[col].isna().all(), col


def test_original_b31g_renames_the_strength_columns(inputs):
    settings = inputs.settings.model_copy(update={
        "strength": inputs.settings.strength.model_copy(update={"method": "original_b31g"})})
    names = list(build_analysis_table(dataclasses.replace(inputs, settings=settings)).columns)
    assert "B31G Burst Pressure (psi)" in names and "B31G Safe Pressure (psi)" in names
    assert "Mod B31G Burst Pressure (psi)" not in names
    assert len(names) == len(NAMES)


def test_report_stage_writes_the_table_and_workbook(inputs, built, tmp_path):  # noqa: F811
    from ili.report import OUTPUT_NAMES, ReportHeader, write_report
    assert ANALYSIS_TABLE_FILE in OUTPUT_NAMES and ANALYSIS_XLSX_FILE in OUTPUT_NAMES
    paths, _ = write_report(inputs, ReportHeader(None, None, None), tmp_path, datetime.now(timezone.utc))
    back = load_analysis_table(paths[ANALYSIS_TABLE_FILE])
    assert list(back.columns) == NAMES and len(back) == len(inputs.classification)
    assert pd.api.types.is_string_dtype(back["Feature ID"])
    ws = openpyxl.load_workbook(paths[ANALYSIS_XLSX_FILE]).active
    assert ws.title == "Classified anomalies"
    assert DRAFT_LABEL in ws["A1"].value
    assert [c.value for c in ws[2]] == NAMES
    assert ws.max_row - 2 == len(back)
    assert ws.freeze_panes == "C3" and ws.auto_filter.ref.startswith("A2:")
    import json
    manifest = json.loads(paths["run_manifest.json"].read_text(encoding="utf-8"))
    assert {ANALYSIS_TABLE_FILE, ANALYSIS_XLSX_FILE} <= set(manifest["outputs"])


def test_missing_file_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        load_analysis_table(tmp_path / ANALYSIS_TABLE_FILE)

"""M12 parity: summary, IMP class comparison and the command line (synthetic pipeline only)."""
import hashlib
import json
import shutil

import openpyxl
import pandas as pd
import pytest
from test_review_flags import plain  # noqa: F401 (module fixture)
from test_parity_read import CMAP, MAP_PATH

from ili import parity
from ili.classify import load_classification_table
from ili.config import DEFAULT_SETTINGS, load_settings
from ili.growth import load_growth_table
from ili.match import load_match_table
from ili.parity import imp_class_comparison, ours_by_feature
from ili.remaining_life import load_remaining_life_table

LABELS = {v: k for k, v in CMAP.anomalies.rate_sources.items()}


@pytest.fixture(scope="module")
def reported(plain, tmp_path_factory):  # noqa: F811
    """A copy of the plain pipeline with classification and the report run, so imp_class.csv exists."""
    import test_review_flags as t
    from ili import classify, report
    data, out0, std, folder0 = plain
    out = tmp_path_factory.mktemp("parity_cli") / "out"
    shutil.copytree(out0, out)
    folder = out / folder0.name
    assert classify.main(t._classify_args(std, folder, data, out)) == 0
    from test_report_cli import _argv
    assert report.main(_argv((data, out, std, folder))) == 0
    return data, out, std, folder


def _workbook(folder, path, *, imp=True):
    ours = ours_by_feature(load_growth_table(folder / "growth.csv"),
                           load_remaining_life_table(folder / "remaining_life.csv"),
                           load_classification_table(folder / "classification.csv"),
                           load_match_table(folder / "match_table.csv"))
    a = CMAP.anomalies
    fields = ["feature_id", "log_dist_ft", "joint_id", "depth_pct", "prev_depth_pct", "rate_mpy", "rate_source",
              "erf", "rl_yr", "vendor_type"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = a.sheet
    ws.append([a.category_column, *[a.columns[f] for f in fields]])

    def cell(v):
        if v is None or (isinstance(v, float) and v != v):
            return None
        return "inf" if v == float("inf") else v

    rows = []
    for r in ours.itertuples(index=False):
        rows.append({"feature_id": r.feature_id, "log_dist_ft": cell(r.log_dist_ft), "joint_id": None,
                     "depth_pct": cell(r.depth_ours), "prev_depth_pct": None, "rate_mpy": cell(r.rate_ours),
                     "rate_source": LABELS.get(r.method_ours), "erf": cell(r.erf_ours), "rl_yr": cell(r.rl_ours),
                     "vendor_type": "Corrosion"})
    pick = [i for i, r in enumerate(rows) if r["rate_mpy"] is not None and r["rate_mpy"] != "inf"][0]
    rows[pick]["rate_mpy"] += 1.0
    changed = rows[pick]["feature_id"]
    gone = [r for r in ours.itertuples(index=False) if r.feature_id != changed and r.category_ours != "removed"][0]
    rows = [r for r in rows if r["feature_id"] != gone.feature_id]
    rows.append({"feature_id": "CL-1", "log_dist_ft": 5.0, "joint_id": None, "depth_pct": 10.0, "prev_depth_pct": None,
                 "rate_mpy": None, "rate_source": None, "erf": None, "rl_yr": None, "vendor_type": "Cluster"})
    for r in rows:
        ws.append(["Volumetric", *[r[f] for f in fields]])
    if imp:
        m = CMAP.imp_summary
        ims = wb.create_sheet(m.sheet)
        ims.append([m.columns[f] for f in ("imp_class", "n_classified", "n_previously_repaired", "n_unrepaired")])
        for c in m.classes:
            ims.append([c, 1, 0, 1])
    wb.save(path)
    return path, changed


def _copy(reported, tmp_path):
    _, out0, _, folder0 = reported
    out = tmp_path / "out"
    shutil.copytree(out0, out)
    return out / folder0.name


def _run(folder, wb, capsys):
    capsys.readouterr()
    code = parity.main(["--comparison", str(folder), "--consultant", str(wb), "--map", str(MAP_PATH),
                        "--settings", str(DEFAULT_SETTINGS)])
    return code, capsys.readouterr()


def test_cli_end_to_end(reported, tmp_path, capsys):
    folder = _copy(reported, tmp_path)
    wb, changed = _workbook(folder, tmp_path / "c.xlsx")
    code, cap = _run(folder, wb, capsys)
    assert code == 0, cap.err
    out = folder / "parity"
    for name in ("parity_by_anomaly.csv", "parity_summary.json", "parity_imp_class.csv"):
        assert (out / name).is_file()
    s = json.loads((out / "parity_summary.json").read_text(encoding="utf-8"))
    ex = s["counts_by_explanation"]
    assert ex["vendor_cluster_box"] == 1 and ex["not_in_their_set"] == 1
    table = pd.read_csv(out / "parity_by_anomaly.csv", dtype=str, keep_default_na=False)
    shared = table[table.side == "both"]
    assert (shared.explanation != "agree").sum() == 1
    assert shared[shared.explanation != "agree"].feature_id.iloc[0] == changed
    assert s["workbook_sha256"] == hashlib.sha256(wb.read_bytes()).hexdigest()
    assert (out / "parity_by_anomaly.csv").read_text(encoding="utf-8").splitlines()[0] == ",".join(parity.PARITY_COLUMNS)
    assert s["label"] in cap.out and s["label"].startswith("DRAFT")
    assert s["imp_class_compared"] is True and set(ex) == set(parity.EXPLANATIONS)
    assert "unexplained" in cap.out
    g = json.loads((folder / "growth_summary.json").read_text(encoding="utf-8"))
    rls = json.loads((folder / "remaining_life_summary.json").read_text(encoding="utf-8"))
    expect = {k: g.get(k) for k in ("governing_profile", "age_basis", "days_per_year", "rate_decimals",
                                    "shrinking_applies_to", "joint_percentile_pool")}
    expect["depth_basis"] = rls.get("depth_basis")
    assert s["method_ours"] == expect and s["method_ours"]["governing_profile"] is not None


def test_cli_missing_growth_exits_2(reported, tmp_path, capsys):
    folder = _copy(reported, tmp_path)
    wb, _ = _workbook(folder, tmp_path / "c.xlsx")
    (folder / "growth.csv").unlink()
    code, cap = _run(folder, wb, capsys)
    assert code == 2 and "growth.csv" in cap.err


def test_cli_without_imp_class_skips_that_table(reported, tmp_path, capsys):
    folder = _copy(reported, tmp_path)
    wb, _ = _workbook(folder, tmp_path / "c.xlsx")
    (folder / "imp_class.csv").unlink()
    code, cap = _run(folder, wb, capsys)
    assert code == 0, cap.err
    assert not (folder / "parity" / "parity_imp_class.csv").exists()
    # a stale file from an earlier run is removed
    (folder / "parity" / "parity_imp_class.csv").write_text("stale", encoding="utf-8")
    code, cap = _run(folder, wb, capsys)
    assert code == 0, cap.err
    assert not (folder / "parity" / "parity_imp_class.csv").exists()
    s = json.loads((folder / "parity" / "parity_summary.json").read_text(encoding="utf-8"))
    assert s["imp_class_compared"] is False and s["imp_class_note"]


def test_cli_bad_workbook_exits_2(reported, tmp_path, capsys):
    folder = _copy(reported, tmp_path)
    wb = tmp_path / "bad.xlsx"
    openpyxl.Workbook().save(wb)
    code, cap = _run(folder, wb, capsys)
    assert code == 2 and "Classified_Anomalies" in cap.err


def test_imp_class_comparison_counts_none_class():
    ours = pd.DataFrame({"imp_class": ["Immediate", "None", "None", "YTC < 5 yrs"],
                         "previously_repaired": ["False", "True", "False", "False"]})
    theirs = pd.DataFrame({"imp_class": ["Immediate", "None"], "n_classified": [1, 3],
                           "n_previously_repaired": [0, 1], "n_unrepaired": [1, 2]}, dtype=object)
    t = imp_class_comparison(ours, theirs, 5)
    assert list(t.imp_class) == ["Immediate", "60-day", "180-day", "YTC < 5 yrs", "None"]
    none = t[t.imp_class == "None"].iloc[0]
    assert (none.n_classified_ours, none.n_previously_repaired_ours, none.n_unrepaired_ours) == (2, 1, 1)
    assert none.n_classified_theirs == 3
    assert t[t.imp_class == "60-day"].iloc[0].n_classified_theirs is None


def test_imp_class_comparison_uses_the_label_in_the_file():
    ours = pd.DataFrame({"imp_class": ["Immediate", "YTC < 3 yrs", "YTC < 3 yrs"],
                         "previously_repaired": ["False", "False", "False"]})
    theirs = pd.DataFrame({"imp_class": ["YTC < 3 yrs"], "n_classified": [2], "n_previously_repaired": [0],
                           "n_unrepaired": [2]}, dtype=object)
    t = imp_class_comparison(ours, theirs, 5)  # settings say 5, the run used 3
    row = t[t.imp_class == "YTC < 3 yrs"].iloc[0]
    assert row.n_classified_ours == 2 and row.n_classified_theirs == 2
    assert "YTC < 5 yrs" not in set(t.imp_class)

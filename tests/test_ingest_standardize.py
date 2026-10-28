from datetime import date

import pandas as pd
import pytest
import yaml
from conftest import MAP_A, MAP_B, assert_matches_oracle

from ili.ingest import VendorMap, load_standardized, load_vendor_map, parse_clock, read_tally, standardize
from ili.schema import STANDARD_COLUMNS


def _ingest(path, map_path, run_id, run_date):
    vmap = load_vendor_map(map_path)
    raw, _ = read_tally(path, vmap)
    return standardize(raw, vmap, run_id, run_date)


def test_vendor_a_matches_oracle(synth):
    result, out = synth
    std, issues = _ingest(out / "tally_prev_vendor_a.csv", MAP_A, "2021_vendorA", date(2021, 6, 15))
    assert issues.empty
    assert list(std.columns) == STANDARD_COLUMNS + ["source_row"]
    assert_matches_oracle(std, result.prev_std)


def test_vendor_b_matches_oracle(synth):
    result, out = synth
    std, issues = _ingest(out / "tally_curr_vendor_b.xlsx", MAP_B, "2026_vendorB", date(2026, 6, 15))
    assert issues.empty
    assert_matches_oracle(std, result.curr_std)


def test_source_rows(synth):
    _, out = synth
    vmap = load_vendor_map(MAP_B)
    raw, _ = read_tally(out / "tally_curr_vendor_b.xlsx", vmap)
    assert raw.source_row.iloc[0] == 4  # header on row 3
    vmap = load_vendor_map(MAP_A)
    raw, _ = read_tally(out / "tally_prev_vendor_a.csv", vmap)
    assert raw.source_row.iloc[0] == 2


def test_excel_ids_stay_integers(synth):
    _, out = synth
    raw, _ = read_tally(out / "tally_curr_vendor_b.xlsx", load_vendor_map(MAP_B))
    gw = raw["Ref GW"][raw["Ref GW"] != ""]
    assert gw.str.fullmatch(r"\d+").all()


def test_run_date_from_title_row(synth):
    _, out = synth
    assert read_tally(out / "tally_curr_vendor_b.xlsx", load_vendor_map(MAP_B))[1] == date(2026, 6, 15)
    assert read_tally(out / "tally_prev_vendor_a.csv", load_vendor_map(MAP_A))[1] is None


def test_parse_clock_formats():
    clock, bad = parse_clock(pd.Series(["12:30", "3:00", "", "13:00", "x"]), "hhmm")
    assert clock.tolist()[:2] == [30, 180] and pd.isna(clock[2])
    assert bad.tolist() == [False, False, False, True, True]
    clock, bad = parse_clock(pd.Series(["90", "360", "359.9", "abc"]), "degrees")
    assert clock.tolist()[:3] == [180, 0, 0] and bad.tolist() == [False, False, False, True]
    clock, bad = parse_clock(pd.Series(["3.5", "12.0", "0"]), "decimal_hours")
    assert clock.tolist() == [210, 0, 0] and not bad.any()


def test_parse_clock_keeps_out_of_range_for_validation():
    clock, bad = parse_clock(pd.Series(["400", "-10", "360.2", "-0.2"]), "degrees")
    assert clock.tolist() == [800, -20, 720, 0] and not bad.any()  # -0.2 deg rounds to 0 min
    clock, _ = parse_clock(pd.Series(["13.0", "-1"]), "decimal_hours")
    assert clock.tolist() == [780, -60]


def test_decimal_hours_twelve_oclock_face():
    clock, bad = parse_clock(pd.Series(["12.5", "12.99", "0.5"]), "decimal_hours")
    assert clock.tolist() == [30, 59, 30] and not bad.any()


def _vendor_a_as_excel(synth, tmp_path, *, percent_depth=False, time_clock=True):
    """Vendor A's tally rewritten as an .xlsx the way a person would build it in Excel."""
    import datetime as dt

    import openpyxl
    import yaml

    _, out = synth
    df = pd.read_csv(out / "tally_prev_vendor_a.csv", dtype=str, keep_default_na=False)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tally"
    ws.append(list(df.columns))
    for r, row in enumerate(df.itertuples(index=False), start=2):
        for c, (col, val) in enumerate(zip(df.columns, row), start=1):
            cell = ws.cell(row=r, column=c)
            if val == "":
                continue
            if col == "O'clock" and time_clock:
                h, m = map(int, val.split(":"))
                cell.value = dt.time(h % 12 if h == 12 else h, m)  # Excel stores 12:30 as 00:30 or 12:30
                cell.number_format = "h:mm"
            elif col == "Depth (%)" and percent_depth:
                cell.value = float(val) / 100
                cell.number_format = "0.0%"
            else:
                try:
                    cell.value = float(val)
                except ValueError:
                    cell.value = val
    path = tmp_path / "vendor_a.xlsx"
    wb.save(path)
    raw_map = yaml.safe_load(MAP_A.read_text(encoding="utf-8"))
    raw_map["file"] = {"format": "excel", "sheet": "Tally", "header_row": 1}
    map_path = tmp_path / "vendor_a_excel.yaml"
    map_path.write_text(yaml.safe_dump(raw_map), encoding="utf-8")
    return path, map_path


def test_excel_percent_formatted_depth(synth, tmp_path):
    path, map_path = _vendor_a_as_excel(synth, tmp_path, percent_depth=True, time_clock=False)
    std, issues = _ingest(path, map_path, "2021_vendorA", date(2021, 6, 15))
    assert issues.empty
    assert_matches_oracle(std, synth[0].prev_std)


def test_excel_time_formatted_clock(synth, tmp_path):
    path, map_path = _vendor_a_as_excel(synth, tmp_path, time_clock=True)
    std, issues = _ingest(path, map_path, "2021_vendorA", date(2021, 6, 15))
    assert issues.empty
    assert_matches_oracle(std, synth[0].prev_std)


def test_csv_source_row_is_physical_line(synth, tmp_path):
    """Blank lines and quoted multi-line cells must not shift source_row."""
    _, out = synth
    df = pd.read_csv(out / "tally_prev_vendor_a.csv", dtype=str, keep_default_na=False)
    df.loc[2, "Comments"] = "first line\nsecond line"
    ml = df.index[df.Event == "Metal Loss"][0]
    target = df.loc[ml, "Feature No"]
    text = df.to_csv(index=False, lineterminator="\n").split("\n")
    text.insert(6, "")  # blank line after physical line 6
    path = tmp_path / "quirky.csv"
    path.write_text("\n".join(text), encoding="utf-8", newline="")
    expected_line = next(i for i, line in enumerate(text, start=1) if line.startswith(target + ","))
    raw, _ = read_tally(path, load_vendor_map(MAP_A))
    assert raw.loc[raw["Feature No"] == target, "source_row"].iloc[0] == expected_line
    assert raw.loc[raw["Feature No"] == df.loc[2, "Feature No"], "Comments"].iloc[0] == "first line\nsecond line"


def _raw_b(**overrides):
    row = {
        "Log Dist [m]": "1", "Feature": "ML-CORR", "Feature ID": "B-000001", "Ref GW": "1",
        "Rel. Dist [m]": "0.5", "Joint Len [m]": "12", "WT [mm]": "6.35", "Peak Depth [%]": "20",
        "Dent Depth [mm]": "", "Length [mm]": "25.4", "Width [mm]": "50.8", "Orientation [deg]": "90",
        "Location": "EXT", "Seam": "0", "Dent ML": "", "Cluster": "", "Lat": "32", "Lon": "-84",
        "Elev [m]": "100", "Comment": "",
    }
    row.update(overrides)
    df = pd.DataFrame([row])
    df["source_row"] = 4
    return df


def test_unit_conversions():
    std, issues = standardize(_raw_b(), load_vendor_map(MAP_B), "r", date(2026, 1, 1))
    assert issues.empty
    row = std.iloc[0]
    assert row.log_dist_ft == pytest.approx(3.28084, abs=1e-5)
    assert row.wt_nom_in == pytest.approx(0.25)
    assert row.length_in == pytest.approx(1.0) and row.width_in == pytest.approx(2.0)
    assert row.elev_ft == pytest.approx(328.084, abs=1e-3)
    assert row.clock_min == 180 and row.seam_interaction is False and row.dent_metal_loss is None


def test_bom_csv_header(synth, tmp_path):
    _, out = synth
    text = (out / "tally_prev_vendor_a.csv").read_text(encoding="utf-8")
    bom = tmp_path / "bom.csv"
    bom.write_text(text, encoding="utf-8-sig")
    vmap = load_vendor_map(MAP_A)
    raw, _ = read_tally(bom, vmap)
    assert "Feature No" in raw.columns
    std, issues = standardize(raw, vmap, "2021_vendorA", date(2021, 6, 15))
    assert issues.empty and std.feature_id.iloc[0] == "A-000001"


def test_vocab_whitespace_case():
    std, issues = standardize(_raw_b(Feature=" ml-corr ", Location=" ext"), load_vendor_map(MAP_B), "r",
                              date(2026, 1, 1))
    assert issues.empty
    assert std.feature_type.iloc[0] == "metal_loss" and std.surface.iloc[0] == "EXT"


def test_missing_mapped_column_raises(synth, tmp_path):
    _, out = synth
    df = pd.read_csv(out / "tally_prev_vendor_a.csv", dtype=str).drop(columns=["O'clock"])
    path = tmp_path / "missing.csv"
    df.to_csv(path, index=False)
    with pytest.raises(ValueError, match="O'clock"):
        read_tally(path, load_vendor_map(MAP_A))


def _origin_map(**ml_origins):
    raw = yaml.safe_load(MAP_B.read_text(encoding="utf-8"))
    raw["feature_types"] = {"ML-CORR": "metal_loss", "ML-MILL": "metal_loss", "DENT": "dent"}
    raw.pop("ml_origins", None)
    if ml_origins:
        raw["ml_origins"] = ml_origins
    return VendorMap.model_validate(raw)


def _origins(vmap, *codes):
    raw = pd.concat([_raw_b(Feature=c) for c in codes], ignore_index=True)
    std, issues = standardize(raw, vmap, "r", date(2026, 1, 1))
    assert issues.empty
    return list(std.ml_origin)


def test_ml_origin_from_vendor_code():
    vmap = _origin_map(**{"ML-CORR": "corrosion", "ML-MILL": "manufacturing"})
    assert _origins(vmap, "ML-CORR", "ML-MILL", "DENT") == ["corrosion", "manufacturing", None]


def test_ml_origin_defaults_unknown():
    assert _origins(_origin_map(**{"ML-CORR": "corrosion"}), "ML-CORR", "ML-MILL", "DENT") == [
        "corrosion", "unknown", None]
    assert _origins(_origin_map(), "ML-CORR", "DENT") == ["unknown", None]


def test_load_standardized_without_ml_origin_column(tmp_path):
    std, _ = standardize(_raw_b(Feature="ML-CORR"), _origin_map(), "r", date(2026, 1, 1))
    dent, _ = standardize(_raw_b(Feature="DENT"), _origin_map(), "r", date(2026, 1, 1))
    path = tmp_path / "old.csv"
    pd.concat([std, dent]).drop(columns=["ml_origin"]).to_csv(path, index=False)
    loaded = load_standardized(path)
    assert list(loaded.ml_origin) == ["unknown", None]


# --- vendor-map options for real vendor exports (ROSEN, L930) ---------------------------------------------------


def _map_b(**changes):
    raw = yaml.safe_load(MAP_B.read_text(encoding="utf-8"))
    raw.update(changes)
    return VendorMap.model_validate(raw)


def test_day_fraction_clock_is_an_excel_time_on_a_12_hour_face():
    minutes, bad = parse_clock(pd.Series(["0.125", "0.484", "0.5", "0.625", "", "x", "0.6"]), "day_fraction")
    # 0.125 day = 3:00 = 180 min; 0.484 day = 11:37; 0.5 = 12:00 wraps to 0; 0.625 = 15:00 is past 13 h: kept for
    # validation to report, like a decimal hour of 15
    assert minutes.tolist()[:4] == [180, 697, 0, 900]
    assert pd.isna(minutes[4]) and bad.tolist() == [False, False, False, False, False, True, False]


def test_negate_flips_a_numeric_field():
    vmap = _map_b(negate=["dist_from_us_gw_ft"])
    std, issues = standardize(_raw_b(**{"Rel. Dist [m]": "-0.5"}), vmap, "r", date(2026, 1, 1))
    assert issues.empty and std.dist_from_us_gw_ft.iloc[0] == pytest.approx(0.5 / 0.3048)


def test_negate_and_fill_down_accept_only_known_fields(tmp_path):
    for option in ("negate", "fill_down"):
        with pytest.raises(ValueError, match="nope"):
            _map_b(**{option: ["nope"]})
    with pytest.raises(ValueError, match="numeric"):
        _map_b(negate=["comments"])


def test_fill_down_carries_joint_fields_from_the_weld_rows():
    vmap = _map_b(fill_down=["us_gw_id", "joint_length_ft", "wt_nom_in"])
    raw = pd.concat([_raw_b(Feature="ML-CORR", **{"Ref GW": "", "Joint Len [m]": "", "WT [mm]": ""}),
                     _raw_b(Feature="GWLD", **{"Ref GW": "10", "Joint Len [m]": "12", "WT [mm]": "9.525"}),
                     _raw_b(Feature="ML-CORR", **{"Ref GW": "", "Joint Len [m]": "", "WT [mm]": ""}),
                     _raw_b(Feature="ML-CORR", **{"Ref GW": "20", "Joint Len [m]": "", "WT [mm]": ""})],
                    ignore_index=True)
    std, _ = standardize(raw, vmap, "r", date(2026, 1, 1))
    assert std.us_gw_id.tolist() == [None, "10", "10", "20"]  # nothing above the first weld to carry
    assert std.wt_nom_in.iloc[2] == pytest.approx(0.375) and pd.isna(std.wt_nom_in.iloc[0])


def test_feature_id_from_row_names_each_row_by_its_source_row():
    raw = yaml.safe_load(MAP_B.read_text(encoding="utf-8"))
    raw["columns"].pop("feature_id")
    vmap = VendorMap.model_validate({**raw, "feature_id_from_row": True})
    std, issues = standardize(_raw_b(), vmap, "r", date(2026, 1, 1))
    assert issues.empty and std.feature_id.iloc[0] == "row-4"
    with pytest.raises(ValueError, match="feature_id_from_row"):
        VendorMap.model_validate({**yaml.safe_load(MAP_B.read_text(encoding="utf-8")), "feature_id_from_row": True})


def test_xls_tally_with_a_units_row(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    path = tmp_path / "tally.xls"
    book = xlwt.Workbook()
    sheet = book.add_sheet("Pipe Tally")
    time_style = xlwt.easyxf(num_format_str="h:mm")
    for c, h in enumerate(["J. no.", "log dist.", "event", "o'clock", "depth"]):
        sheet.write(2, c, h)
        sheet.write(3, c, "[units]")
    sheet.write(4, 0, 2150.0)
    sheet.write(4, 1, 12.5)
    sheet.write(4, 2, "metal loss corrosion")
    sheet.write(4, 3, 0.125, time_style)
    sheet.write(4, 4, 20.0)
    book.save(str(path))
    from ili.ingest import FileSpec
    vmap = VendorMap.model_construct(file=FileSpec(format="xls", header_row=3, skip_rows=1),
                                     columns={"us_gw_id": "J. no.", "log_dist_ft": "log dist.",
                                              "feature_type": "event", "clock_min": "o'clock", "depth_pct": "depth"},
                                     run_date=None, vendor="t")
    raw, _ = read_tally(path, vmap)
    assert len(raw) == 1 and raw.source_row.tolist() == [5]
    row = raw.iloc[0]
    assert (row["J. no."], row["log dist."], row["o'clock"], row["depth"]) == ("2150", "12.5", "3:00", "20")
